"""Bounded XLSX inspection with hardened XML and no formula evaluation."""
import io
from pathlib import Path, PurePosixPath
import posixpath
import re
from zipfile import ZipFile, BadZipFile

from defusedxml import ElementTree as SafeXML
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple

from mathmodel_paper_studio.inputs.profiling import (
    ImportLimits, checked_table, read_bounded, read_source_bounded,
)

_NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
_RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def _check_reference(reference: str, limits: ImportLimits) -> None:
    try:
        left, top, right, bottom = range_boundaries(reference)
        if (None in (left, top, right, bottom) or left < 1 or top < 1
                or right < left or bottom < top or right > limits.max_columns
                or bottom > limits.max_rows + 1 or right * bottom > limits.max_cells):
            raise ValueError('XLSX 单元格范围超过导入上限。')
    except (TypeError, ValueError) as exc:
        raise ValueError('XLSX 单元格范围无效或超过导入上限。') from exc


def _text_length(element, limits: ImportLimits) -> int:
    # openpyxl matches children by localname, including foreign namespaces.
    # Reject unsupported forms before counting only the standard text paths.
    if any(not node.tag.startswith(_NS) for node in element.iter()):
        raise ValueError('XLSX 富文本包含不受支持的命名空间。')
    supported = {'t', 'r', 'rPh', 'phoneticPr'}
    for parent in (element, *element.findall(f'{_NS}r')):
        allowed = supported if parent is element else {'rPr', 't'}
        if any(node.tag not in {f'{_NS}{tag}' for tag in allowed} for node in parent):
            raise ValueError('XLSX 富文本包含不受支持的子结构。')
    # Count incrementally, without building the concatenated value.
    length = 0
    for node in (*element.findall(f'{_NS}t'), *element.findall(f'{_NS}r/{_NS}t')):
        length += len(node.text or '')
        if length > limits.max_cell_chars:
            raise ValueError('XLSX 富文本合并后的单元格内容超过长度上限。')
    return length


def _shared_lengths(roots: dict, limits: ImportLimits) -> tuple[int, ...]:
    for root in roots.values():
        for tag in ('si', 'is'):
            for element in root.iter(f'{_NS}{tag}'):
                _text_length(element, limits)
    content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml'
    parts = [node.attrib.get('PartName', '') for node in roots['[Content_Types].xml']
             if node.attrib.get('ContentType') == content_type]
    if not parts:
        return ()
    if len(parts) != 1 or not parts[0].startswith('/'):
        raise ValueError('XLSX 共享字符串表声明无效或重复。')
    root = roots.get(parts[0][1:])
    if root is None or root.tag != f'{_NS}sst':
        raise ValueError('XLSX 共享字符串表不存在或格式无效。')
    return tuple(_text_length(element, limits) for element in root.iter(f'{_NS}si'))


def _check_cell_text(cell, shared: tuple[int, ...], limits: ImportLimits) -> None:
    if cell.attrib.get('t') != 's':
        return
    values = cell.findall(f'{_NS}v')
    index = values[0].text if len(values) == 1 else None
    if index is None or re.fullmatch(r'[0-9]+', index) is None:
        raise ValueError('XLSX 共享字符串索引必须为非负整数。')
    number = int(index)
    if number >= len(shared):
        raise ValueError('XLSX 共享字符串索引超出有效范围。')
    if shared[number] > limits.max_cell_chars:
        raise ValueError('XLSX 共享字符串内容超过单元格长度上限。')


def _check_sheet(root, limits: ImportLimits, shared: tuple[int, ...]) -> None:
    dimension = root.find(f'{_NS}dimension')
    if dimension is not None:
        _check_reference(dimension.attrib.get('ref', ''), limits)
    rows = root.findall(f'{_NS}sheetData/{_NS}row')
    if len(rows) > limits.max_rows + 1:
        raise ValueError('XLSX 行数超过上限。')
    cells, previous_row = 0, 0
    for row in rows:
        row_text = row.attrib.get('r', '')
        if re.fullmatch(r'[1-9][0-9]*', row_text) is None:
            raise ValueError('XLSX 行坐标必须为正整数。')
        row_number = int(row_text)
        if not previous_row < row_number <= limits.max_rows + 1:
            raise ValueError('XLSX 行坐标重复、乱序或超过上限。')
        previous_row, previous_column = row_number, 0
        for cell in row.findall(f'{_NS}c'):
            cells += 1
            reference = cell.attrib.get('r', '')
            if re.fullmatch(r'[A-Z]+[1-9][0-9]*', reference) is None:
                raise ValueError('XLSX 单元格必须使用单个规范坐标。')
            _check_reference(reference, limits)
            cell_row, column = coordinate_to_tuple(reference)
            if cell_row != row_number or column <= previous_column:
                raise ValueError('XLSX 单元格行归属错误或列坐标重复、乱序。')
            previous_column = column
            _check_cell_text(cell, shared, limits)
            if cells > limits.max_cells:
                raise ValueError('XLSX 单元格总数超过上限。')
    for merge in root.findall(f'{_NS}mergeCells/{_NS}mergeCell'):
        _check_reference(merge.attrib.get('ref', ''), limits)

def _inspect(content: bytes, limits: ImportLimits) -> dict:
    roots = {}
    with ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > limits.max_zip_members:
            raise ValueError('XLSX ZIP 成员数量超过上限。')
        names = set()
        total = 0
        for info in infos:
            name = info.filename
            parts = PurePosixPath(name).parts
            if (not parts or name.startswith('/') or '\\' in name or ':' in name
                    or '..' in parts or name.casefold() in names or info.flag_bits & 1):
                raise ValueError('XLSX ZIP 路径重复、无效或加密。')
            names.add(name.casefold())
            total += info.file_size
            if (info.file_size > limits.max_zip_member_bytes or info.compress_size > limits.max_file_bytes
                    or total > limits.max_zip_total_bytes):
                raise ValueError('XLSX ZIP 解压大小超过上限。')
            lower = name.casefold()
            if ('externallinks/' in lower or 'vbaproject' in lower
                    or lower.startswith(('xl/embeddings/', 'xl/activex/'))):
                raise ValueError('不支持带外部链接、宏或嵌入对象的 XLSX。')
        for info in infos:
            if not info.filename.casefold().endswith(('.xml', '.rels')):
                continue
            with archive.open(info) as stream:
                data = stream.read(limits.max_zip_member_bytes + 1)
            if len(data) > limits.max_zip_member_bytes:
                raise ValueError('XLSX ZIP 成员超过上限。')
            root = SafeXML.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            roots[info.filename] = root
            for node in root.iter():
                if (node.attrib.get('TargetMode', '').casefold() == 'external'
                        or 'macroenabled' in node.attrib.get('ContentType', '').casefold()):
                    raise ValueError('XLSX 包含外部链接或宏。')
                if node.text and len(node.text) > limits.max_cell_chars:
                    raise ValueError('XLSX 文本或公式超过单元格长度上限。')
    shared = _shared_lengths(roots, limits)
    for name, root in roots.items():
        if name.startswith('xl/worksheets/') and name.endswith('.xml'):
            _check_sheet(root, limits, shared)
    return roots


def _sheet_paths(roots: dict) -> dict[str, str]:
    relationships = {node.attrib['Id']: node.attrib['Target']
                     for node in roots['xl/_rels/workbook.xml.rels']}
    paths = {}
    for sheet in roots['xl/workbook.xml'].findall(f'{_NS}sheets/{_NS}sheet'):
        target = relationships[sheet.attrib[_RID]]
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        if path not in roots or not path.startswith('xl/worksheets/'):
            raise ValueError('XLSX 工作表引用无效。')
        paths[sheet.attrib['name']] = path
    return paths


def xlsx_sheet_names(path: Path | bytes, limits: ImportLimits = ImportLimits()) -> tuple[str, ...]:
    """Inspect an imported XLSX source read-only, before it is copied to app storage."""
    content = read_source_bounded(path, limits) if isinstance(path, Path) else read_bounded(path, limits)
    try:
        return tuple(_sheet_paths(_inspect(content, limits)))
    except ValueError:
        raise
    except (BadZipFile, DefusedXmlException, OSError, KeyError, TypeError, OverflowError, SyntaxError, EOFError, RuntimeError) as exc:
        raise ValueError('XLSX 文件无效、XML 不安全或结构不受支持。') from exc


def read_xlsx(path: Path | bytes, *, sheet: str | None = None, limits: ImportLimits = ImportLimits()):
    content = read_bounded(path, limits)
    workbook = None
    try:
        roots = _inspect(content, limits)
        paths = _sheet_paths(roots)
        names = tuple(paths)
        if sheet not in paths:
            if sheet is None and len(names) == 1:
                sheet = names[0]
            else:
                raise ValueError('请明确选择工作表；可用工作表：' + '、'.join(names))
        root = roots[paths[sheet]]
        merged = tuple(node.attrib['ref'] for node in root.findall(f'{_NS}mergeCells/{_NS}mergeCell'))
        # Preflight already inspected every XML member with defusedxml before
        # the library sees bytes. No source path, external link, or VBA is loaded.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False,
                                 keep_links=False, keep_vba=False)
        worksheet = workbook[sheet]
        # Derive actual coordinates as dimension metadata may understate data.
        cells = root.findall(f'{_NS}sheetData/{_NS}row/{_NS}c')
        coordinates = tuple(coordinate_to_tuple(cell.attrib['r']) for cell in cells)
        height = max((row for row, _ in coordinates), default=0)
        width = max((column for _, column in coordinates), default=0)
        if height * width > limits.max_cells:
            raise ValueError('XLSX 单元格总数超过上限。')
        rows = []
        formulas = []
        headers = ()
        for index, row in enumerate(worksheet.iter_rows(min_row=1, max_row=height, max_col=width)):
            values = tuple(cell.value for cell in row)
            formulas.extend(cell.coordinate for cell in row if cell.data_type == 'f')
            if index == 0:
                headers = values
            else:
                rows.append(values)
        return checked_table(headers, rows, limits, sheet_name=sheet, sheet_names=names,
                             formulas=tuple(formulas), merged_ranges=merged,
                             notes=('xml=defusedxml', 'formulas=literal;not-recalculated', 'external-links=disabled'))
    except ValueError:
        raise
    except (BadZipFile, DefusedXmlException, OSError, KeyError, TypeError, OverflowError, SyntaxError, EOFError, RuntimeError) as exc:
        raise ValueError('XLSX 文件无效、XML 不安全或结构不受支持。') from exc
    finally:
        if workbook is not None:
            workbook.close()
