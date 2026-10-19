"""Versioned OOXML writer that changes only explicitly mapped template anchors."""
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
import hashlib
import os
from pathlib import Path
import re
from uuid import uuid4
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException

from mathmodel_paper_studio.bootstrap.paths import validate_managed_path
from mathmodel_paper_studio.paper.standards import get_default_paper_standard
from mathmodel_paper_studio.paper.template_inspector import TemplateInspection, inspect_template
from mathmodel_paper_studio.paper.template_mapping import TemplateMapping, TemplateMappingError

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
_M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
_R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
_PKG_REL = '{http://schemas.openxmlformats.org/package/2006/relationships}'
_CT = '{http://schemas.openxmlformats.org/package/2006/content-types}'
_WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
_PIC = '{http://schemas.openxmlformats.org/drawingml/2006/picture}'
_XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'
_IMAGE_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'
_MAX_PACKAGE_BYTES = 64 * 1024 * 1024
_MAX_IMAGE_BYTES = 16 * 1024 * 1024
_MAX_TABLE_ROWS = 10000
_MAX_TEXT_CHARS = 100000
_EVIDENCE_ID = re.compile(r'^[a-z][a-z0-9_-]{2,63}$')

for _prefix, _uri in (
    ('w', _W[1:-1]), ('r', _R[1:-1]),
    ('wp', _WP[1:-1]), ('a', _A[1:-1]), ('pic', _PIC[1:-1]),
    ('m', _M[1:-1]),
    ('pr', 'http://schemas.openxmlformats.org/package/2006/relationships'),
    ('ct', 'http://schemas.openxmlformats.org/package/2006/content-types'),
):
    ET.register_namespace(_prefix, _uri)


class DocxWriteError(ValueError):
    """The draft cannot be safely applied to the selected template mapping."""


def _text(value: str, *, allow_empty: bool = False) -> str:
    if not isinstance(value, str) or len(value) > _MAX_TEXT_CHARS:
        raise DocxWriteError('Document text is empty, invalid, or exceeds its size limit.')
    # Model and imported table text can contain XML 1.0-forbidden code points.
    # Remove those before writing OOXML so one bad character cannot pause the
    # paper stage after the Agent has already returned a valid draft.
    compatible = ''.join(
        character for character in value
        if ord(character) in (0x09, 0x0A, 0x0D)
        or 0x20 <= ord(character) <= 0xD7FF
        or 0xE000 <= ord(character) <= 0xFFFD
        or 0x10000 <= ord(character) <= 0x10FFFF
    )
    if not allow_empty and not compatible.strip():
        raise DocxWriteError('Document text is empty, invalid, or exceeds its size limit.')
    return compatible


def _evidence_ids(values) -> tuple[str, ...]:
    if not isinstance(values, (tuple, list)) or len(values) > 64:
        raise DocxWriteError('Evidence ID collection is invalid or too large.')
    result = tuple(values)
    if any(not isinstance(item, str) or not _EVIDENCE_ID.fullmatch(item) for item in result):
        raise DocxWriteError('Evidence ID is invalid.')
    if len(set(result)) != len(result):
        raise DocxWriteError('Evidence IDs cannot be duplicated.')
    return result


@dataclass(frozen=True)
class SubsectionDraft:
    title: str
    paragraphs: tuple[str, ...]
    evidence_ids: tuple[str, ...] = ()
    pending_reviews: tuple[str, ...] = ()

    def __post_init__(self):
        object.__setattr__(self, 'title', _text(self.title))
        if not isinstance(self.paragraphs, (tuple, list)) or not self.paragraphs:
            raise DocxWriteError('Subsection requires at least one paragraph.')
        object.__setattr__(self, 'paragraphs', tuple(_text(item) for item in self.paragraphs))
        object.__setattr__(self, 'evidence_ids', _evidence_ids(self.evidence_ids))
        pending = tuple(_text(item) for item in self.pending_reviews)
        if len(pending) > 32:
            raise DocxWriteError('Too many subsection review markers.')
        object.__setattr__(self, 'pending_reviews', pending)


@dataclass(frozen=True)
class TextSection:
    section_key: str
    paragraphs: tuple[str, ...]
    evidence_ids: tuple[str, ...] = ()
    pending_reviews: tuple[str, ...] = ()
    subsections: tuple[SubsectionDraft, ...] = ()

    def __post_init__(self):
        if not isinstance(self.paragraphs, (tuple, list)) or not self.paragraphs:
            raise DocxWriteError('Text section requires at least one paragraph.')
        paragraphs = tuple(_text(item) for item in self.paragraphs)
        evidence = _evidence_ids(self.evidence_ids)
        pending = tuple(_text(item) for item in self.pending_reviews)
        if len(pending) > 32:
            raise DocxWriteError('Too many pending review markers.')
        if not isinstance(self.subsections, (tuple, list)) or len(self.subsections) > 32:
            raise DocxWriteError('Subsections are invalid or exceed the section limit.')
        subsections = tuple(self.subsections)
        if any(not isinstance(item, SubsectionDraft) for item in subsections):
            raise DocxWriteError('Text section contains an invalid subsection.')
        object.__setattr__(self, 'paragraphs', paragraphs)
        object.__setattr__(self, 'evidence_ids', evidence)
        object.__setattr__(self, 'pending_reviews', pending)
        object.__setattr__(self, 'subsections', subsections)


@dataclass(frozen=True)
class TableDraft:
    section_key: str
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    evidence_ids: tuple[str, ...] = ()
    pending_reviews: tuple[str, ...] = ()
    caption: str = ''

    def __post_init__(self):
        headers = tuple(_text(item) for item in self.headers)
        rows = tuple(tuple(_text(cell, allow_empty=True) for cell in row) for row in self.rows)
        if not headers or len(headers) > 128 or len(rows) > _MAX_TABLE_ROWS:
            raise DocxWriteError('Table dimensions are empty or exceed limits.')
        if any(len(row) != len(headers) for row in rows):
            raise DocxWriteError('Table rows must have the same column count as the header.')
        evidence = _evidence_ids(self.evidence_ids)
        pending = tuple(_text(item) for item in self.pending_reviews)
        caption = _text(self.caption, allow_empty=True)
        object.__setattr__(self, 'headers', headers)
        object.__setattr__(self, 'rows', rows)
        object.__setattr__(self, 'evidence_ids', evidence)
        object.__setattr__(self, 'pending_reviews', pending)
        object.__setattr__(self, 'caption', caption)


@dataclass(frozen=True)
class FigureDraft:
    section_key: str
    image_bytes: bytes
    image_format: str
    caption: str
    evidence_ids: tuple[str, ...] = ()
    pending_reviews: tuple[str, ...] = ()

    def __post_init__(self):
        _text(self.caption)
        if not isinstance(self.image_bytes, bytes) or not 8 <= len(self.image_bytes) <= _MAX_IMAGE_BYTES:
            raise DocxWriteError('Figure image must be a bounded PNG or JPEG byte sequence.')
        if self.image_format not in ('png', 'jpeg'):
            raise DocxWriteError('Only PNG and JPEG figures are supported.')
        if self.image_format == 'png' and not self.image_bytes.startswith(b'\x89PNG\r\n\x1a\n'):
            raise DocxWriteError('PNG figure signature is invalid.')
        if self.image_format == 'jpeg' and not (
                self.image_bytes.startswith(b'\xff\xd8\xff') and self.image_bytes.endswith(b'\xff\xd9')):
            raise DocxWriteError('JPEG figure signature is invalid.')
        object.__setattr__(self, 'evidence_ids', _evidence_ids(self.evidence_ids))
        object.__setattr__(self, 'pending_reviews', tuple(_text(item) for item in self.pending_reviews))


@dataclass(frozen=True)
class EquationDraft:
    section_key: str
    expression: str
    description: str
    evidence_ids: tuple[str, ...] = ()
    pending_reviews: tuple[str, ...] = ()

    def __post_init__(self):
        if not isinstance(self.section_key, str) or not re.fullmatch(r'[a-z][a-z0-9_-]{0,63}', self.section_key):
            raise DocxWriteError('Equation section key is invalid.')
        object.__setattr__(self, 'expression', _text(self.expression))
        object.__setattr__(self, 'description', _text(self.description))
        object.__setattr__(self, 'evidence_ids', _evidence_ids(self.evidence_ids))
        pending = tuple(_text(item) for item in self.pending_reviews)
        if len(pending) > 32:
            raise DocxWriteError('Too many pending equation review markers.')
        object.__setattr__(self, 'pending_reviews', pending)


@dataclass(frozen=True)
class PaperDraft:
    sections: tuple[TextSection, ...] = ()
    tables: tuple[TableDraft, ...] = ()
    figures: tuple[FigureDraft, ...] = ()
    pending_references: tuple[str, ...] = ()
    equations: tuple[EquationDraft, ...] = ()
    title: str = '数学建模论文'
    keywords: tuple[str, ...] = ()

    def __post_init__(self):
        for name in ('sections', 'tables', 'figures', 'pending_references', 'equations', 'keywords'):
            value = getattr(self, name)
            if not isinstance(value, (tuple, list)):
                raise DocxWriteError(f'{name} must be a list or tuple.')
            object.__setattr__(self, name, tuple(value))
        if any(not isinstance(item, TextSection) for item in self.sections):
            raise DocxWriteError('Paper sections contain an invalid value.')
        if any(not isinstance(item, TableDraft) for item in self.tables):
            raise DocxWriteError('Paper tables contain an invalid value.')
        if any(not isinstance(item, FigureDraft) for item in self.figures):
            raise DocxWriteError('Paper figures contain an invalid value.')
        if any(not isinstance(item, EquationDraft) for item in self.equations):
            raise DocxWriteError('Paper equations contain an invalid value.')
        keys = [item.section_key for item in (*self.sections, *self.tables, *self.figures)]
        if len(set(keys)) != len(keys):
            raise DocxWriteError('Each draft block requires a distinct mapping key.')
        if len(self.pending_references) > 256:
            raise DocxWriteError('Too many pending references.')
        object.__setattr__(self, 'pending_references', tuple(_text(item) for item in self.pending_references))
        object.__setattr__(self, 'title', _text(self.title))
        keywords = tuple(_text(item) for item in self.keywords)
        if len(keywords) > 32:
            raise DocxWriteError('Too many paper keywords.')
        object.__setattr__(self, 'keywords', keywords)
        section_keys = {item.section_key for item in self.sections}
        if any(item.section_key not in section_keys for item in self.equations):
            raise DocxWriteError('Equation must refer to a declared paper section.')


@dataclass(frozen=True)
class DocxWriteResult:
    path: Path
    sha256: str
    source_sha256: str
    version: int
    modified_parts: tuple[str, ...]


_GENERIC_SECTION_TITLES = {
    'abstract': '摘要', 'problem_restatement': '问题重述', 'problem_analysis': '问题分析',
    'assumptions': '模型假设', 'notation': '符号说明', 'model': '模型建立',
    'solution': '模型求解与分析', 'results': '结果与讨论', 'verification': '模型检验',
    'sensitivity_analysis': '灵敏度分析', 'robustness_check': '稳健性分析',
    'computational_complexity': '算法复杂度', 'conclusion': '结论与模型评价',
    'references': '参考文献', 'appendix': '附录',
}

_GENERIC_SECTION_ORDER = (
    'abstract', 'problem_restatement', 'problem_analysis', 'assumptions', 'notation',
    'model', 'solution', 'results', 'verification', 'sensitivity_analysis',
    'robustness_check', 'computational_complexity', 'conclusion', 'references', 'appendix',
)

_CHINESE_NUMERALS = ('一', '二', '三', '四', '五', '六', '七', '八', '九', '十',
                     '十一', '十二', '十三', '十四', '十五', '十六', '十七', '十八', '十九', '二十')


def _generic_section_title(section_key: str) -> str:
    return _GENERIC_SECTION_TITLES.get(section_key, section_key.replace('_', ' '))


def _generic_figure_anchor_title(index: int) -> str:
    return f'结果图 {index}'


def _append_generic_artifact_anchors(document, draft: PaperDraft) -> None:
    for table_draft in draft.tables:
        document.add_paragraph('{{table_caption:' + table_draft.section_key + '}}', style='Caption')
        table = document.add_table(rows=1, cols=len(table_draft.headers))
        table.style = 'Table Grid'
        for column_index, cell in enumerate(table.rows[0].cells):
            cell.text = '{{table:' + table_draft.section_key + ':' + str(column_index) + '}}'
    for figure_index, _figure in enumerate(draft.figures, start=1):
        document.add_heading(_generic_figure_anchor_title(figure_index), level=3)


def _numbered_section_title(section_key: str) -> str:
    title = _generic_section_title(section_key)
    if section_key == 'abstract':
        return title
    try:
        index = _GENERIC_SECTION_ORDER.index(section_key)
    except ValueError:
        return title
    number = index if index < len(_CHINESE_NUMERALS) else index + 1
    return f'{_CHINESE_NUMERALS[number - 1]}、{title}'


def _section_decimal_number(section_key: str) -> int:
    try:
        return max(1, _GENERIC_SECTION_ORDER.index(section_key))
    except ValueError:
        return 1


def _set_east_asia_font(style, font_name: str) -> None:
    style.font.name = font_name
    properties = style._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        properties.insert(0, fonts)
    fonts.set(qn('w:eastAsia'), font_name)


def _generic_document_bytes(draft: PaperDraft) -> bytes:
    document = Document()
    style_defaults = get_default_paper_standard().style_defaults
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    margin_cm = float(style_defaults['margins'].split()[0])
    section.top_margin = section.bottom_margin = Cm(margin_cm)
    section.left_margin = section.right_margin = Cm(margin_cm)
    normal = document.styles['Normal']
    _set_east_asia_font(normal, style_defaults['body_font'])
    normal.font.size = Pt(float(style_defaults['body_size_pt']))
    normal.paragraph_format.line_spacing = float(style_defaults['body_line_spacing'])
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    title_style = document.styles['Title']
    _set_east_asia_font(title_style, style_defaults['title_font'])
    title_style.font.size = Pt(float(style_defaults['title_size_pt']))
    title_style.font.bold = True
    title_style.paragraph_format.alignment = 1
    title_style.paragraph_format.space_after = Pt(14)
    for level in (1, 2, 3):
        heading = document.styles[f'Heading {level}']
        _set_east_asia_font(heading, style_defaults['heading_font'])
        heading.font.size = Pt(float(style_defaults[f'heading{level}_size_pt']))
        heading.font.bold = True
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.first_line_indent = Cm(0)
    caption = document.styles['Caption']
    _set_east_asia_font(caption, style_defaults['body_font'])
    caption.font.size = Pt(float(style_defaults['caption_size_pt']))
    caption.paragraph_format.alignment = 1
    caption.paragraph_format.first_line_indent = Cm(0)

    document.add_paragraph('{{title}}', style='Title')
    ordered_sections = sorted(draft.sections, key=lambda item: (
        _GENERIC_SECTION_ORDER.index(item.section_key)
        if item.section_key in _GENERIC_SECTION_ORDER else len(_GENERIC_SECTION_ORDER),
        item.section_key,
    ))
    has_abstract = any(item.section_key == 'abstract' for item in ordered_sections)
    if not has_abstract:
        document.add_paragraph('{{keywords}}', style='Caption')
    section_keys = {item.section_key for item in ordered_sections}
    artifact_anchor = next(
        (key for key in ('results', 'solution', 'model', 'problem_analysis')
         if key in section_keys),
        None,
    )
    fallback_before = next(
        (item.section_key for item in ordered_sections
         if item.section_key in ('conclusion', 'references', 'appendix')),
        None,
    ) if artifact_anchor is None else None
    artifacts_added = False
    for item in ordered_sections:
        if fallback_before == item.section_key and (draft.tables or draft.figures):
            document.add_heading(_numbered_section_title('results') + '（待补正文）', level=1)
            document.add_paragraph('结果正文待补充；请结合下方已验证数据补充分析。')
            _append_generic_artifact_anchors(document, draft)
            artifacts_added = True
        document.add_heading(_numbered_section_title(item.section_key), level=1)
        document.add_paragraph('{{section:' + item.section_key + '}}')
        if item.section_key == 'abstract':
            document.add_paragraph('{{keywords}}', style='Caption')
        if item.section_key == artifact_anchor and (draft.tables or draft.figures):
            _append_generic_artifact_anchors(document, draft)
            artifacts_added = True
    if (draft.tables or draft.figures) and not artifacts_added:
        document.add_heading(_numbered_section_title('results') + '（待补正文）', level=1)
        document.add_paragraph('结果正文待补充；请结合下方已验证数据补充分析。')
        _append_generic_artifact_anchors(document, draft)

    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    run = footer.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    result = OxmlElement('w:t')
    result.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend((begin, instruction, separate, result, end))

    package = BytesIO()
    try:
        document.save(package)
    except (OSError, ValueError, TypeError) as exc:
        raise DocxWriteError('Built-in Word document could not be assembled.') from exc
    content = package.getvalue()
    if not content or len(content) > _MAX_PACKAGE_BYTES:
        raise DocxWriteError('Built-in Word document is empty or exceeds the package limit.')
    return content


def _save_generic_template(output_directory: Path, content: bytes) -> Path:
    template_path = None
    for version in range(1, 10000):
        destination = validate_managed_path(
            output_directory / f'.builtin-template-v{version:04d}.docx')
        if destination.exists():
            continue
        temporary = validate_managed_path(
            output_directory / f'.builtin-template-v{version:04d}.{uuid4().hex}.tmp')
        try:
            with temporary.open('xb') as stream:
                stream.write(content)
                stream.flush()
                os.fsync(stream.fileno())
            validate_managed_path(temporary)
            validate_managed_path(destination)
            if destination.exists():
                temporary.unlink()
                continue
            temporary.rename(destination)
            template_path = validate_managed_path(destination)
            break
        except FileExistsError:
            if temporary.exists():
                temporary.unlink()
        except BaseException:
            if temporary.exists():
                validate_managed_path(temporary)
                temporary.unlink()
            raise
    if template_path is None:
        raise DocxWriteError('No free built-in Word template version remains.')
    return template_path


def _generic_template_mapping(template_path: Path, draft: PaperDraft):
    inspection = inspect_template(template_path)
    headings = {item.label: item.anchor_id for item in inspection.anchors if item.kind == 'heading'}
    tables = iter(item for item in inspection.anchors if item.kind == 'table')
    selections = {}
    try:
        for item in draft.sections:
            selections[item.section_key] = headings[_numbered_section_title(item.section_key)]
        for index, item in enumerate(draft.figures, start=1):
            selections[item.section_key] = headings[_generic_figure_anchor_title(index)]
        for item in draft.tables:
            selections[item.section_key] = next(tables).anchor_id
    except (StopIteration, KeyError) as exc:
        raise DocxWriteError('Built-in Word anchors do not match the generated draft.') from exc
    mapping = TemplateMapping.from_selection(inspection, selections)
    mapping.validate_for(inspection)
    return template_path, inspection, mapping


def write_generic_template(output_directory: Path, draft: PaperDraft):
    """Create a fresh editable Word skeleton and its app-owned anchor mapping."""
    if not isinstance(draft, PaperDraft):
        raise DocxWriteError('Paper draft type is invalid.')
    output_directory = validate_managed_path(Path(output_directory))
    if not output_directory.is_dir():
        raise DocxWriteError('Generic Word output directory is unavailable.')
    template_path = _save_generic_template(output_directory, _generic_document_bytes(draft))
    return _generic_template_mapping(template_path, draft)


def _safe_parse(content: bytes, part: str):
    try:
        SafeET.fromstring(content, forbid_dtd=True, forbid_entities=True, forbid_external=True)
        parser = ET.XMLParser(target=ET.TreeBuilder(insert_comments=True, insert_pis=True))
        return ET.fromstring(content, parser=parser)
    except (DefusedXmlException, ET.ParseError, ValueError) as exc:
        raise DocxWriteError(f'Template XML is invalid: {part}') from exc


def _serialize(root, original: bytes) -> bytes:
    namespaces = []
    seen = set()
    for _, pair in ET.iterparse(BytesIO(original), events=('start-ns',)):
        prefix, uri = pair
        if pair not in seen:
            seen.add(pair)
            namespaces.append((prefix, uri))
            try:
                ET.register_namespace(prefix, uri)
            except ValueError:
                pass
    data = ET.tostring(root, encoding='utf-8', xml_declaration=True, short_empty_elements=True)
    declaration_end = data.find(b'?>')
    root_start = data.find(b'<', declaration_end + 2)
    root_end = data.find(b'>', root_start)
    for prefix, uri in namespaces:
        token = b'xmlns=' if not prefix else b'xmlns:' + prefix.encode('utf-8') + b'='
        if token in data[root_start:root_end]:
            continue
        escaped = uri.replace('&', '&amp;').replace('"', '&quot;')
        attribute = (b' xmlns="' if not prefix else b' xmlns:' + prefix.encode('utf-8') + b'="')
        addition = attribute + escaped.encode('utf-8') + b'"'
        data = data[:root_end] + addition + data[root_end:]
        root_end += len(addition)
    return data


def _all_parents(root):
    return {child: parent for parent in root.iter() for child in list(parent)}


def _set_paragraph_text(paragraph, text: str) -> None:
    ppr = paragraph.find(_W + 'pPr')
    first_run = paragraph.find(_W + 'r')
    rpr = None if first_run is None else first_run.find(_W + 'rPr')
    retained = [child for child in list(paragraph)
                if child.tag in (_W + 'pPr', _W + 'bookmarkStart', _W + 'bookmarkEnd',
                                 _W + 'commentRangeStart', _W + 'commentRangeEnd', _W + 'proofErr')]
    paragraph[:] = []
    if ppr is not None:
        paragraph.append(ppr)
    run = ET.Element(_W + 'r')
    if rpr is not None:
        run.append(deepcopy(rpr))
    node = ET.SubElement(run, _W + 't')
    if text[:1].isspace() or text[-1:].isspace():
        node.set(_XML_SPACE, 'preserve')
    node.text = text
    children = [item for item in retained if item.tag != _W + 'pPr']
    closing_marks = {_W + 'bookmarkEnd', _W + 'commentRangeEnd'}
    insertion = next((i for i, item in enumerate(children) if item.tag in closing_marks), len(children))
    children.insert(insertion, run)
    paragraph[:] = ([ppr] if ppr is not None else []) + children


def _paragraph(text: str):
    result = ET.Element(_W + 'p')
    _set_paragraph_text(result, text)
    return result


def _styled_paragraph(text: str, style_id: str):
    result = _paragraph(text)
    ppr = ET.Element(_W + 'pPr')
    style = ET.SubElement(ppr, _W + 'pStyle')
    style.set(_W + 'val', style_id)
    result.insert(0, ppr)
    return result


def _math_paragraph(expression: str):
    symbols = {
        '\\alpha': 'α', '\\beta': 'β', '\\gamma': 'γ', '\\delta': 'δ',
        '\\Delta': 'Δ', '\\theta': 'θ', '\\lambda': 'λ', '\\mu': 'μ',
        '\\sigma': 'σ', '\\omega': 'ω', '\\pi': 'π', '\\sum': '∑',
        '\\prod': '∏', '\\cdot': '·', '\\times': '×', '\\leq': '≤',
        '\\geq': '≥', '\\neq': '≠', '\\infty': '∞',
    }

    def add_run(parent, value):
        run = ET.SubElement(parent, _M + 'r')
        text = ET.SubElement(run, _M + 't')
        text.text = value
        if value[:1].isspace() or value[-1:].isspace():
            text.set(_XML_SPACE, 'preserve')

    def read_script(value, offset):
        if offset >= len(value):
            return '', offset, False
        if value[offset] == '{':
            end = value.find('}', offset + 1)
            if end < 0:
                return value[offset + 1:], len(value), False
            raw = value[offset + 1:end]
            supported = '{' not in raw and '}' not in raw
            return raw, end + 1, supported
        if value[offset] == '\\':
            match = re.match(r'\\[A-Za-z]+', value[offset:])
            if match:
                macro = match.group(0)
                return symbols.get(macro, macro), offset + len(macro), macro in symbols
        return value[offset], offset + 1, True

    paragraph = ET.Element(_W + 'p')
    paragraph_properties = ET.SubElement(paragraph, _W + 'pPr')
    ET.SubElement(paragraph_properties, _W + 'jc', {_W + 'val': 'center'})
    math = ET.SubElement(paragraph, _M + 'oMath')
    index = 0
    supported = True
    while index < len(expression):
        if expression[index] == '\\':
            match = re.match(r'\\[A-Za-z]+', expression[index:])
            if match:
                macro = match.group(0)
                atom = symbols.get(macro, macro)
                supported = supported and macro in symbols
                index += len(macro)
            else:
                atom = expression[index]
                index += 1
                supported = False
        elif expression[index].isascii() and expression[index].isalnum():
            match = re.match(r'[A-Za-z0-9]+', expression[index:])
            atom = match.group(0)
            index += len(atom)
        else:
            atom = expression[index]
            index += 1

        subscript = None
        superscript = None
        while index < len(expression) and expression[index] in '_^':
            marker = expression[index]
            value, index, script_supported = read_script(expression, index + 1)
            supported = supported and script_supported
            if not value:
                supported = False
                break
            if marker == '_' and subscript is None:
                subscript = value
            elif marker == '^' and superscript is None:
                superscript = value
            else:
                supported = False
                add_run(math, marker + value)
        if subscript is None and superscript is None:
            add_run(math, atom)
            continue
        if subscript is not None and superscript is not None:
            structured = ET.SubElement(math, _M + 'sSubSup')
            elements = (('e', atom), ('sub', subscript), ('sup', superscript))
        elif subscript is not None:
            structured = ET.SubElement(math, _M + 'sSub')
            elements = (('e', atom), ('sub', subscript))
        else:
            structured = ET.SubElement(math, _M + 'sSup')
            elements = (('e', atom), ('sup', superscript))
        for name, value in elements:
            container = ET.SubElement(structured, _M + name)
            add_run(container, value)
    return paragraph, supported


def _after(parent, target, elements) -> None:
    index = list(parent).index(target) + 1
    for offset, element in enumerate(elements):
        parent.insert(index + offset, element)


def _section_placeholder(parent, heading, token, heading_nodes):
    siblings = list(parent)
    for sibling in siblings[siblings.index(heading) + 1:]:
        if sibling in heading_nodes or sibling.tag == _W + 'sectPr':
            break
        if sibling.tag == _W + 'tbl':
            break
        if sibling.tag != _W + 'p':
            continue
        if token in ''.join(node.text or '' for node in sibling.iter(_W + 't')):
            return sibling
    return None


def _markers(evidence_ids, pending_reviews=()):
    result = []
    if evidence_ids:
        result.append('证据：' + ', '.join(evidence_ids))
    else:
        result.append('待人工核验：内容尚未关联证据')
    result.extend('待人工核验：' + item for item in pending_reviews)
    return result


def _replace_table(table, draft: TableDraft):
    rows = table.findall(_W + 'tr')
    if not rows:
        raise DocxWriteError('Mapped table has no header row to preserve.')
    grid = table.find(_W + 'tblGrid')
    column_count = len(grid.findall(_W + 'gridCol')) if grid is not None else len(rows[0].findall(_W + 'tc'))
    if column_count != len(draft.headers) or len(rows[0].findall(_W + 'tc')) != column_count:
        raise DocxWriteError('Mapped table columns do not match the confirmed draft schema.')
    _set_row(rows[0], draft.headers)
    data_template = deepcopy(rows[1] if len(rows) > 1 else rows[0])
    for row in rows[1:]:
        table.remove(row)
    header_flag = data_template.find(f'{_W}trPr/{_W}tblHeader')
    if header_flag is not None:
        data_template.find(_W + 'trPr').remove(header_flag)
    for values in draft.rows:
        row = deepcopy(data_template)
        _set_row(row, values)
        table.append(row)


def _set_row(row, values):
    cells = row.findall(_W + 'tc')
    if len(cells) != len(values):
        raise DocxWriteError('Mapped table row has an unexpected cell count.')
    for cell, value in zip(cells, values):
        paragraph = cell.find(_W + 'p')
        if paragraph is None:
            paragraph = ET.SubElement(cell, _W + 'p')
        _set_paragraph_text(paragraph, value)
        for extra in cell.findall(_W + 'p')[1:]:
            cell.remove(extra)


def _drawing_paragraph(relationship_id: str, figure: FigureDraft, drawing_id: int):
    paragraph = ET.Element(_W + 'p')
    run = ET.SubElement(paragraph, _W + 'r')
    drawing = ET.SubElement(run, _W + 'drawing')
    inline = ET.SubElement(drawing, _WP + 'inline', {'distT': '0', 'distB': '0', 'distL': '0', 'distR': '0'})
    ET.SubElement(inline, _WP + 'extent', {'cx': '4572000', 'cy': '3429000'})
    evidence = ', '.join(figure.evidence_ids) or 'not linked'
    ET.SubElement(inline, _WP + 'docPr', {'id': str(drawing_id), 'name': f'figure-{figure.section_key}',
                                          'descr': f'Evidence: {evidence}'})
    frame = ET.SubElement(inline, _WP + 'cNvGraphicFramePr')
    ET.SubElement(frame, _A + 'graphicFrameLocks', {'noChangeAspect': '1'})
    graphic = ET.SubElement(ET.SubElement(inline, _A + 'graphic'), _A + 'graphicData',
                            {'uri': 'http://schemas.openxmlformats.org/drawingml/2006/picture'})
    picture = ET.SubElement(graphic, _PIC + 'pic')
    nv = ET.SubElement(picture, _PIC + 'nvPicPr')
    ET.SubElement(nv, _PIC + 'cNvPr', {'id': str(drawing_id), 'name': f'figure-{figure.section_key}',
                                       'descr': f'Evidence: {evidence}'})
    ET.SubElement(nv, _PIC + 'cNvPicPr')
    fill = ET.SubElement(picture, _PIC + 'blipFill')
    ET.SubElement(fill, _A + 'blip', {_R + 'embed': relationship_id})
    ET.SubElement(ET.SubElement(fill, _A + 'stretch'), _A + 'fillRect')
    shape = ET.SubElement(picture, _PIC + 'spPr')
    transform = ET.SubElement(shape, _A + 'xfrm')
    ET.SubElement(transform, _A + 'off', {'x': '0', 'y': '0'})
    ET.SubElement(transform, _A + 'ext', {'cx': '4572000', 'cy': '3429000'})
    geometry = ET.SubElement(shape, _A + 'prstGeom', {'prst': 'rect'})
    ET.SubElement(geometry, _A + 'avLst')
    return paragraph


def _add_content_type(root, extension: str, mime_type: str) -> None:
    for item in root.findall(_CT + 'Default'):
        if item.get('Extension', '').casefold() == extension.casefold():
            if item.get('ContentType') != mime_type:
                raise DocxWriteError('Template declares a conflicting image content type.')
            return
    ET.SubElement(root, _CT + 'Default', {'Extension': extension, 'ContentType': mime_type})


def _replace_placeholder(paragraph, token: str, replacement: str) -> bool:
    current = ''.join(node.text or '' for node in paragraph.iter(_W + 't'))
    if token not in current:
        return False
    _set_paragraph_text(paragraph, current.replace(token, replacement))
    return True


def _binding_for(mapping, key, kind):
    binding = next((item for item in mapping.bindings if item.section_key == key), None)
    if binding is None or binding.kind != kind:
        raise DocxWriteError(f'Draft block {key!r} has no confirmed {kind} mapping.')
    return binding


def write_paper(template_path: Path, output_directory: Path, inspection: TemplateInspection,
                mapping: TemplateMapping, draft: PaperDraft) -> DocxWriteResult:
    """Write a new versioned DOCX. The source template is never opened for writing."""
    if not isinstance(inspection, TemplateInspection) or not isinstance(mapping, TemplateMapping):
        raise DocxWriteError('Template inspection and mapping are required.')
    if not isinstance(draft, PaperDraft):
        raise DocxWriteError('Paper draft type is invalid.')
    mapping.validate_for(inspection)
    try:
        template_path = validate_managed_path(Path(template_path))
        output_directory = validate_managed_path(Path(output_directory))
        if not template_path.is_file() or not output_directory.is_dir():
            raise DocxWriteError('Template or paper output directory is unavailable.')
        with template_path.open('rb') as stream:
            source_bytes = stream.read(_MAX_PACKAGE_BYTES + 1)
        if not source_bytes or len(source_bytes) > _MAX_PACKAGE_BYTES:
            raise DocxWriteError('Template package is empty or exceeds the 64 MiB limit.')
    except (OSError, ValueError, TypeError) as exc:
        if isinstance(exc, DocxWriteError):
            raise
        raise DocxWriteError('Template or output path cannot be safely read.') from exc
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    if source_hash != inspection.template_sha256:
        raise TemplateMappingError('Template changed after inspection; confirm mapping again.')
    try:
        with ZipFile(BytesIO(source_bytes)) as archive:
            original_infos = archive.infolist()
            parts = {info.filename: archive.read(info) for info in original_infos}
            archive_comment = archive.comment
    except (BadZipFile, OSError, RuntimeError, ValueError) as exc:
        raise DocxWriteError('Template package could not be copied safely.') from exc

    document_original = parts['word/document.xml']
    rels_original = parts['word/_rels/document.xml.rels']
    types_original = parts['[Content_Types].xml']
    document = _safe_parse(document_original, 'word/document.xml')
    relationships = _safe_parse(rels_original, 'word/_rels/document.xml.rels')
    content_types = _safe_parse(types_original, '[Content_Types].xml')
    body = document.find(_W + 'body')
    if body is None:
        raise DocxWriteError('Template document body is missing.')
    parent_map = _all_parents(document)
    paragraphs = list(body.iter(_W + 'p'))
    tables = list(body.iter(_W + 'tbl'))
    by_id = {item.anchor_id: item for item in inspection.anchors}
    anchor_nodes = {}
    for anchor in inspection.anchors:
        nodes = paragraphs if anchor.kind == 'heading' else tables
        if anchor.index >= len(nodes):
            raise DocxWriteError('Mapped anchor index is outside the document.')
        anchor_nodes[anchor.anchor_id] = nodes[anchor.index]
    heading_nodes = {anchor_nodes[item.anchor_id] for item in inspection.anchors if item.kind == 'heading'}
    for binding in mapping.bindings:
        if binding.anchor_id not in by_id or binding.anchor_id not in anchor_nodes:
            raise TemplateMappingError('Mapped anchor no longer exists in the inspected template.')

    body_paragraphs = list(body.iter(_W + 'p'))
    title_token = '{{title}}'
    keywords_token = '{{keywords}}'
    title_replaced = any(_replace_placeholder(item, title_token, draft.title) for item in body_paragraphs)
    keywords_text = '关键词：' + ('；'.join(draft.keywords) if draft.keywords else '待补充')
    keywords_replaced = any(_replace_placeholder(item, keywords_token, keywords_text)
                            for item in body_paragraphs)
    if not title_replaced:
        body.insert(0, _styled_paragraph(draft.title, 'Title'))
    if not keywords_replaced and not any(item.section_key == 'abstract' for item in draft.sections):
        body.insert(1 if not title_replaced else 0, _styled_paragraph(keywords_text, 'Caption'))

    section_tails = {}
    for section in draft.sections:
        binding = _binding_for(mapping, section.section_key, 'heading')
        heading = anchor_nodes[binding.anchor_id]
        parent = parent_map.get(heading)
        if parent is None:
            raise DocxWriteError('Mapped heading has no writable parent.')
        token = '{{section:' + section.section_key + '}}'
        inserted = []
        placeholder = _section_placeholder(parent, heading, token, heading_nodes)
        if placeholder is not None:
            _replace_placeholder(placeholder, token, section.paragraphs[0])
        if placeholder is not None:
            inserted.append(placeholder)
            inserted.extend(_paragraph(value) for value in section.paragraphs[1:])
            _after(parent, placeholder, inserted[1:])
        else:
            inserted = [_paragraph(value) for value in section.paragraphs]
            _after(parent, heading, inserted)
        marker_paragraphs = [_paragraph(text) for text in _markers(section.evidence_ids, section.pending_reviews)]
        _after(parent, inserted[-1], marker_paragraphs)
        tail = marker_paragraphs[-1] if marker_paragraphs else inserted[-1]
        if section.section_key == 'abstract' and not keywords_replaced:
            keyword_paragraph = _styled_paragraph(keywords_text, 'Caption')
            _after(parent, tail, [keyword_paragraph])
            tail = keyword_paragraph
        chapter_number = _section_decimal_number(section.section_key)
        for subsection_number, subsection in enumerate(section.subsections, start=1):
            subsection_heading = _styled_paragraph(
                f'{chapter_number}.{subsection_number} {subsection.title}', 'Heading2')
            subsection_body = [_paragraph(value) for value in subsection.paragraphs]
            subsection_markers = [_paragraph(text) for text in _markers(
                subsection.evidence_ids, subsection.pending_reviews)]
            additions = [subsection_heading, *subsection_body, *subsection_markers]
            _after(parent, tail, additions)
            tail = additions[-1]
        section_tails[section.section_key] = tail
        parent_map = _all_parents(document)

    for equation_number, equation in enumerate(draft.equations, start=1):
        binding = _binding_for(mapping, equation.section_key, 'heading')
        heading = anchor_nodes[binding.anchor_id]
        parent = _all_parents(document).get(heading)
        if parent is None:
            raise DocxWriteError('Mapped equation section has no writable parent.')
        tail = section_tails.get(equation.section_key, heading)
        caption = _styled_paragraph(f'公式（{equation_number}）：{equation.description}', 'Caption')
        equation_paragraph, syntax_supported = _math_paragraph(equation.expression)
        extras = [equation_paragraph, caption]
        if not syntax_supported:
            extras.append(_paragraph('待人工核验：公式保留为可编辑线性数学文本，部分语法未转换为二维数学结构。'))
        extras.extend(_paragraph(text) for text in _markers(equation.evidence_ids, equation.pending_reviews))
        _after(parent, tail, extras)
        section_tails[equation.section_key] = extras[-1]

    for table_number, table_draft in enumerate(draft.tables, start=1):
        binding = _binding_for(mapping, table_draft.section_key, 'table')
        table = anchor_nodes[binding.anchor_id]
        _replace_table(table, table_draft)
        parent = parent_map.get(table)
        if parent is None:
            raise DocxWriteError('Mapped table has no writable parent.')
        caption_text = table_draft.caption or _generic_section_title(table_draft.section_key)
        caption_value = f'表{table_number}：{caption_text}'
        caption_token = '{{table_caption:' + table_draft.section_key + '}}'
        caption_replaced = any(_replace_placeholder(item, caption_token, caption_value)
                               for item in body.iter(_W + 'p'))
        if not caption_replaced:
            parent.insert(list(parent).index(table), _styled_paragraph(caption_value, 'Caption'))
            parent_map = _all_parents(document)
            parent = parent_map.get(table)
        markers = [_paragraph(text) for text in _markers(table_draft.evidence_ids, table_draft.pending_reviews)]
        _after(parent, table, markers)
        parent_map = _all_parents(document)

    existing_ids = {item.get('Id') for item in relationships.findall(_PKG_REL + 'Relationship')}
    existing_draw_ids = [int(node.get('id')) for node in document.iter(_WP + 'docPr')
                         if (node.get('id') or '').isdigit()]
    drawing_id = max(existing_draw_ids, default=0) + 1
    figure_number = 0
    relationship_number = 0
    for figure in draft.figures:
        binding = _binding_for(mapping, figure.section_key, 'heading')
        heading = anchor_nodes[binding.anchor_id]
        parent = _all_parents(document).get(heading)
        if parent is None:
            raise DocxWriteError('Mapped figure heading has no writable parent.')
        figure_number += 1
        extension = 'png' if figure.image_format == 'png' else 'jpg'
        mime_type = 'image/png' if extension == 'png' else 'image/jpeg'
        media_name = f'word/media/mmps-figure-{figure_number:04d}.{extension}'
        if media_name in parts:
            raise DocxWriteError('Figure media path already exists in the template.')
        relationship_number += 1
        relationship_id = f'rIdMMPS{relationship_number:04d}'
        while relationship_id in existing_ids:
            relationship_number += 1
            relationship_id = f'rIdMMPS{relationship_number:04d}'
        existing_ids.add(relationship_id)
        ET.SubElement(relationships, _PKG_REL + 'Relationship', {
            'Id': relationship_id, 'Type': _IMAGE_REL, 'Target': 'media/' + Path(media_name).name,
        })
        _add_content_type(content_types, extension, mime_type)
        parts[media_name] = figure.image_bytes
        figure_paragraph = _drawing_paragraph(relationship_id, figure, drawing_id)
        drawing_id += 1
        extras = [figure_paragraph, _styled_paragraph(f'图{figure_number}：{figure.caption}', 'Caption')]
        extras.extend(_paragraph(text) for text in _markers(figure.evidence_ids, figure.pending_reviews))
        _after(parent, heading, extras)

    if not any(item.section_key == 'references' for item in draft.sections):
        for pending in draft.pending_references:
            body.insert(max(0, len(list(body)) - (1 if body.find(_W + 'sectPr') is not None else 0)),
                        _paragraph('待补文献：' + pending))

    parts['word/document.xml'] = _serialize(document, document_original)
    parts['word/_rels/document.xml.rels'] = _serialize(relationships, rels_original)
    parts['[Content_Types].xml'] = _serialize(content_types, types_original)
    modified = ('[Content_Types].xml', 'word/_rels/document.xml.rels', 'word/document.xml')
    output_bytes = BytesIO()
    try:
        with ZipFile(BytesIO(source_bytes)) as source, ZipFile(output_bytes, 'w', ZIP_DEFLATED) as output:
            output.comment = archive_comment
            for info in original_infos:
                output.writestr(info, parts[info.filename])
            for name in sorted(set(parts) - set(source.namelist())):
                output.writestr(name, parts[name], compress_type=ZIP_DEFLATED)
    except (BadZipFile, OSError, RuntimeError, ValueError) as exc:
        raise DocxWriteError('Generated DOCX package could not be assembled.') from exc
    package_bytes = output_bytes.getvalue()
    if len(package_bytes) > _MAX_PACKAGE_BYTES:
        raise DocxWriteError('Generated DOCX exceeds the 64 MiB package limit.')

    for version in range(1, 10000):
        destination = validate_managed_path(output_directory / f'paper-v{version:04d}.docx')
        if destination.exists():
            continue
        temporary = validate_managed_path(output_directory / f'.paper-v{version:04d}.{uuid4().hex}.tmp')
        try:
            with temporary.open('xb') as stream:
                stream.write(package_bytes)
                stream.flush()
                os.fsync(stream.fileno())
            validate_managed_path(temporary)
            validate_managed_path(destination)
            if destination.exists():
                temporary.unlink()
                continue
            temporary.rename(destination)
            validate_managed_path(destination)
            return DocxWriteResult(
                destination, hashlib.sha256(package_bytes).hexdigest(), source_hash, version,
                tuple(sorted(set(modified) | {name for name in parts if name not in {info.filename for info in original_infos}})),
            )
        except FileExistsError:
            if temporary.exists():
                temporary.unlink()
            continue
        except BaseException:
            if temporary.exists():
                validate_managed_path(temporary)
                temporary.unlink()
            raise
    raise DocxWriteError('No free versioned paper output name remains.')
