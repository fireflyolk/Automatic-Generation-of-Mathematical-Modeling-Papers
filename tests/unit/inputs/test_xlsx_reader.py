from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from mathmodel_paper_studio.inputs.xlsx_reader import read_xlsx, xlsx_sheet_names
from mathmodel_paper_studio.inputs.profiling import ImportLimits, profile_table
ROOT = Path(r'F:\MathModelPaperStudio\tmp\mathmodel-paper-studio-20261015')
class XlsxTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        self.tmp = tempfile.TemporaryDirectory(dir=ROOT)
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'data.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '数据'
        ws.append(['name','value'])
        ws.append(['甲',2])
        ws.append(['乙','=2+2'])
        ws.merge_cells('A4:B4')
        wb.create_sheet('其他').append(['x'])
        wb.save(self.path)
        wb.close()
    def test_reads_verified_byte_snapshot_with_size_limit(self):
        content = self.path.read_bytes()
        table = read_xlsx(content, sheet='数据')
        self.assertEqual('甲', table.rows[0][0])
        with self.assertRaises(ValueError):
            read_xlsx(content, sheet='数据', limits=ImportLimits(max_file_bytes=1))
    def test_explicit_sheet_formula_merge_no_save(self):
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError,'数据.*其他'):
            read_xlsx(self.path)
        table = read_xlsx(self.path,sheet='数据')
        self.assertEqual(('数据','其他'),table.sheet_names)
        self.assertEqual(('B3',),table.formulas)
        self.assertEqual(('A4:B4',),table.merged_ranges)
        self.assertEqual('=2+2',table.rows[1][1])
        self.assertEqual(before,self.path.read_bytes())
        self.assertEqual(profile_table(table),profile_table(read_xlsx(self.path,sheet='数据')))
    def test_sheet_names_use_the_same_safety_preflight(self):
        self.assertEqual(('数据', '其他'), xlsx_sheet_names(self.path))
        self.assertEqual(('数据', '其他'), xlsx_sheet_names(self.path.read_bytes()))
        with self.assertRaises(ValueError):
            xlsx_sheet_names(self.path, ImportLimits(max_file_bytes=1))
        self.path.write_bytes(b'not zip')
        with self.assertRaises(ValueError):
            xlsx_sheet_names(self.path)

    def test_sheet_names_allow_read_only_sources_on_other_local_drives(self):
        source = Path(r'C:\Users\example\Downloads\data.xlsx')
        content = self.path.read_bytes()
        with patch('mathmodel_paper_studio.inputs.xlsx_reader.read_source_bounded',
                   return_value=content) as read_source, \
                patch('mathmodel_paper_studio.inputs.xlsx_reader.read_bounded',
                      side_effect=AssertionError('managed storage reader used for source file')):
            self.assertEqual(('数据', '其他'), xlsx_sheet_names(source))
        read_source.assert_called_once_with(source, ImportLimits())
    def test_zip_and_parse_limits(self):
        for limit in (ImportLimits(max_file_bytes=1),ImportLimits(max_zip_member_bytes=1),
                      ImportLimits(max_zip_total_bytes=1),ImportLimits(max_rows=1),
                      ImportLimits(max_columns=1),ImportLimits(max_cell_chars=1)):
            with self.subTest(limit=limit), self.assertRaises(ValueError):
                read_xlsx(self.path,sheet='数据',limits=limit)
        self.path.write_bytes(b'not zip')
        with self.assertRaises(ValueError):
            read_xlsx(self.path)
    def test_external_macro_entity_and_duplicate_members_rejected(self):
        original = self.path.read_bytes()
        for name,content in [('xl/externalLinks/externalLink1.xml',b'<x/>'),('xl/vbaProject.bin',b'vba'),
                             ('xl/worksheets/sheet1.xml',b'<!DOCTYPE x [<!ENTITY e "bad">]><x>&e;</x>')]:
            self.path.write_bytes(original)
            with ZipFile(self.path,'a') as archive:
                if name == 'xl/worksheets/sheet1.xml':
                    with self.assertWarnsRegex(UserWarning, 'Duplicate name'):
                        archive.writestr(name, content)
                else:
                    archive.writestr(name, content)
            with self.assertRaises(ValueError):
                read_xlsx(self.path,sheet='数据')

    def test_missing_sheet_and_dimension_bomb(self):
        with self.assertRaisesRegex(ValueError, '数据.*其他'):
            read_xlsx(self.path, sheet='missing')
        with ZipFile(self.path) as archive:
            items = [(item, archive.read(item.filename)) for item in archive.infolist()]
        with ZipFile(self.path, 'w') as archive:
            for item, content in items:
                if item.filename == 'xl/worksheets/sheet1.xml':
                    content = content.replace(b'A1:B4', b'A1:XFD1048576')
                archive.writestr(item, content)
        with self.assertRaises(ValueError):
            read_xlsx(self.path, sheet='数据')

    def test_entity_without_duplicate_and_external_relationship_refused(self):
        original = self.path.read_bytes()
        for target in ('entity', 'external'):
            self.path.write_bytes(original)
            with ZipFile(self.path) as archive:
                items = [(item, archive.read(item.filename)) for item in archive.infolist()]
            with ZipFile(self.path, 'w') as archive:
                for item, content in items:
                    if target == 'entity' and item.filename == 'xl/worksheets/sheet1.xml':
                        content = b'<!DOCTYPE x [<!ENTITY e "bad">]><x>&e;</x>'
                    if target == 'external' and item.filename == '_rels/.rels':
                        content = content.replace(b'</Relationships>', b'<Relationship Id="bad" Type="bad" Target="https://example.invalid" TargetMode="External"/></Relationships>')
                    archive.writestr(item, content)
            with self.subTest(target=target), self.assertRaises(ValueError):
                read_xlsx(self.path, sheet='数据')

    def test_zip_path_traversal_and_case_duplicate_rejected(self):
        original = self.path.read_bytes()
        for name in ('../escape.xml', 'XL/WORKSHEETS/SHEET1.XML'):
            self.path.write_bytes(original)
            with ZipFile(self.path, 'a') as archive:
                archive.writestr(name, b'<x/>')
            with self.subTest(name=name), self.assertRaises(ValueError):
                read_xlsx(self.path, sheet='数据')

    def test_profile_date_bool_and_empty_xlsx_columns(self):
        from datetime import datetime
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(['date', 'flag', 'empty'])
        ws.append([datetime(2026, 1, 2), True, None])
        ws.append([datetime(2026, 1, 3), False, None])
        wb.save(self.path)
        wb.close()
        profile = profile_table(read_xlsx(self.path))
        self.assertEqual(['date', 'boolean', 'empty'], [c['type'] for c in profile['columns']])
        self.assertEqual(2, profile['columns'][2]['missing'])

    def test_actual_coordinates_bypass_small_dimensions(self):
        with ZipFile(self.path) as archive:
            items = [(item, archive.read(item.filename)) for item in archive.infolist()]
        with ZipFile(self.path, 'w') as archive:
            for item, content in items:
                if item.filename == 'xl/worksheets/sheet1.xml':
                    content = content.replace(b'A1:B4', b'A1:A1').replace(b'r="B3"', b'r="XFD3"')
                archive.writestr(item, content)
        with self.assertRaises(ValueError):
            read_xlsx(self.path, sheet='数据')

    def test_member_and_total_cell_count_limits(self):
        for limits in (ImportLimits(max_zip_members=1), ImportLimits(max_cells=3)):
            with self.subTest(limits=limits), self.assertRaises(ValueError):
                read_xlsx(self.path, sheet='数据', limits=limits)

    def test_empty_and_broken_workbook_diagnostic(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.save(self.path)
        wb.close()
        with self.assertRaises(ValueError):
            read_xlsx(self.path)
        for member in ('[Content_Types].xml', 'xl/workbook.xml'):
            with ZipFile(self.path, 'w') as archive:
                archive.writestr(member, b'<x/>')
            with self.subTest(member=member), self.assertRaises(ValueError):
                read_xlsx(self.path)

    def _replace_sheet(self, rows, *, shared=None, shared_part='xl/sharedStrings.xml'):
        if not hasattr(self, '_clean_package'):
            self._clean_package = self.path.read_bytes()
        from io import BytesIO
        with ZipFile(BytesIO(self._clean_package)) as archive:
            items = [(item, archive.read(item.filename)) for item in archive.infolist()]
        ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
        header = '<row r="1"><c r="A1" t="inlineStr"><is><t>a</t></is></c><c r="B1" t="inlineStr"><is><t>b</t></is></c></row>'
        sheet = f'<worksheet xmlns="{ns}"><dimension ref="A1:B4"/><sheetData>{header}{rows}</sheetData></worksheet>'
        with ZipFile(self.path, 'w') as archive:
            for item, content in items:
                if item.filename == 'xl/worksheets/sheet1.xml':
                    content = sheet.encode('utf-8')
                if shared is not None and item.filename == '[Content_Types].xml':
                    part = f'<Override PartName="/{shared_part}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
                    content = content.replace(b'</Types>', part.encode() + b'</Types>')
                if shared is not None and item.filename == 'xl/_rels/workbook.xml.rels':
                    relation = f'<Relationship Id="rIdShared" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="/{shared_part}"/>'
                    content = content.replace(b'</Relationships>', relation.encode() + b'</Relationships>')
                archive.writestr(item, content)
            if shared is not None:
                archive.writestr(shared_part, f'<sst xmlns="{ns}">{shared}</sst>'.encode('utf-8'))

    def _assert_preflight_rejection(self, *, limits=ImportLimits()):
        from unittest.mock import patch
        from mathmodel_paper_studio.inputs import xlsx_reader
        with patch.object(xlsx_reader, 'load_workbook', wraps=xlsx_reader.load_workbook) as load:
            with self.assertRaises(ValueError):
                read_xlsx(self.path, sheet='数据', limits=limits)
            load.assert_not_called()

    def test_duplicate_unordered_or_mismatched_coordinates_rejected_before_load(self):
        cases = {
            'duplicate-rows': '<row r="2"><c r="A2"><v>10</v></c></row><row r="2"><c r="A2"><v>20</v></c></row>',
            'descending-rows': '<row r="3"><c r="A3"><v>10</v></c></row><row r="2"><c r="A2"><v>20</v></c></row>',
            'parent-mismatch': '<row r="3"><c r="A2"><v>20</v></c></row>',
            'duplicate-columns': '<row r="2"><c r="A2"><v>10</v></c><c r="A2"><v>20</v></c></row>',
            'descending-columns': '<row r="2"><c r="B2"><v>10</v></c><c r="A2"><v>20</v></c></row>',
        }
        for case, rows in cases.items():
            with self.subTest(case=case):
                self._replace_sheet(rows)
                self._assert_preflight_rejection()

    def test_cell_coordinates_must_be_single_canonical_coordinates(self):
        for coordinate in ('A2:B2', 'A2:A2', '$A$2', 'a2', 'A02', 'A0', ' A2', 'A2 ', 'A+2', ''):
            with self.subTest(coordinate=coordinate):
                self._replace_sheet(f'<row r="2"><c r="{coordinate}"><v>10</v></c></row>')
                self._assert_preflight_rejection()

    def test_aggregate_inline_and_shared_rich_text_rejected_before_load(self):
        fragments = '<r><t>' + 'x' * 40 + '</t></r><r><t>' + 'y' * 40 + '</t></r>'
        for kind in ('inline', 'shared', 'custom-shared-part'):
            with self.subTest(kind=kind):
                if kind == 'inline':
                    self._replace_sheet(f'<row r="2"><c r="A2" t="inlineStr"><is>{fragments}</is></c></row>')
                else:
                    self._replace_sheet('<row r="2"><c r="A2" t="s"><v>0</v></c></row>',
                                        shared=f'<si>{fragments}</si>',
                                        shared_part='xl/custom-strings.xml' if kind == 'custom-shared-part' else 'xl/sharedStrings.xml')
                self._assert_preflight_rejection(limits=ImportLimits(max_cell_chars=64))

    def test_nonstandard_rich_text_namespaces_and_structure_rejected_before_load(self):
        first, second = 'x' * 40, 'y' * 40
        cases = {
            'unqualified-runs': f'<r xmlns=""><t>{first}</t></r><r xmlns=""><t>{second}</t></r>',
            'foreign-runs': f'<r xmlns="urn:foreign"><t>{first}</t></r><r xmlns="urn:foreign"><t>{second}</t></r>',
            'unqualified-run-text': f'<r><t xmlns="">{first}</t></r><r><t xmlns="">{second}</t></r>',
            'foreign-run-text': f'<r><t xmlns="urn:foreign">{first}</t></r><r><t xmlns="urn:foreign">{second}</t></r>',
            'unqualified-plain-text': f'<t xmlns="">{first}</t><r><t>{second}</t></r>',
            'foreign-plain-text': f'<t xmlns="urn:foreign">{first}</t><r><t>{second}</t></r>',
            'nested-runs': f'<r><r><t>{first}</t></r></r>',
        }
        for kind in ('inline', 'shared'):
            for case, fragments in cases.items():
                with self.subTest(kind=kind, case=case):
                    if kind == 'inline':
                        self._replace_sheet(f'<row r="2"><c r="A2" t="inlineStr"><is>{fragments}</is></c></row>')
                    else:
                        self._replace_sheet('<row r="2"><c r="A2" t="s"><v>0</v></c></row>',
                                            shared=f'<si>{fragments}</si>')
                    before = self.path.read_bytes()
                    self._assert_preflight_rejection(limits=ImportLimits(max_cell_chars=64))
                    self.assertEqual(before, self.path.read_bytes())

    def test_shared_string_indices_are_validated_before_load(self):
        for index in ('-1', '', None, '1', 'x', '0.0'):
            with self.subTest(index=index):
                value = '' if index is None else f'<v>{index}</v>'
                self._replace_sheet(f'<row r="2"><c r="A2" t="s">{value}</c></row>', shared='<si><t>valid</t></si>')
                self._assert_preflight_rejection()
        self._replace_sheet('<row r="2"><c r="A2" t="s"><v>0</v></c></row>')
        self._assert_preflight_rejection()

    def test_valid_rich_and_shared_strings_and_sparse_coordinates_preserved(self):
        self._replace_sheet('<row r="3"><c r="A3" t="s"><v>0</v></c><c r="B3" t="inlineStr"><is><r><t>ab</t></r><r><t>cd</t></r></is></c></row>',
                            shared='<si><r><rPr><b/></rPr><t>12</t></r><r><t>34</t></r></si>', shared_part='xl/custom-strings.xml')
        before = self.path.read_bytes()
        table = read_xlsx(self.path, sheet='数据')
        self.assertEqual(((None, None), ('1234', 'abcd')), table.rows)
        self.assertEqual(before, self.path.read_bytes())
