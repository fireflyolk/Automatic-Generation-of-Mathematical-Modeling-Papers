import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import mathmodel_paper_studio.projects.project_service as project_service_module
from mathmodel_paper_studio.bootstrap.paths import AppPaths
from mathmodel_paper_studio.agents.preflight import Provider
from mathmodel_paper_studio.domain.models import Project, Run, RunStatus
from mathmodel_paper_studio.projects.manifest import latest_manifest
from mathmodel_paper_studio.projects.project_service import ProjectService
from mathmodel_paper_studio.storage.database import Database
from mathmodel_paper_studio.storage.repositories import ProjectRepository, RunRepository
ROOT = Path(r'F:\MathModelPaperStudio\tmp\mathmodel-paper-studio-project-edit-20261018\unit-project-service')
class ProjectServiceTests(unittest.TestCase):
    def setUp(self):
        ROOT.mkdir(parents=True, exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=ROOT)
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.paths = AppPaths(self.base / 'app')
        db = Database(self.paths)
        db.migrate()
        self.repo = ProjectRepository(db)
        self.service = ProjectService(self.paths, self.repo)
    def source(self, name='data.csv', data=b'x,y\n1,2\n'):
        p = self.base / name
        p.write_bytes(data)
        return p

    def test_delete_removes_project_files_and_run_records(self):
        project_id = 'a' * 32
        project = Project(project_id, 'delete me', f'projects/{project_id}')
        self.repo.add(project)
        project_root = self.paths.root / project.relative_path
        (project_root / '00_inputs').mkdir(parents=True)
        (project_root / '00_inputs/problem.txt').write_text('question', encoding='utf-8')
        run_id = 'b' * 32
        runs = RunRepository(Database(self.paths))
        runs.add(Run(run_id, project_id, Provider.CODEX, RunStatus.COMPLETED,
                     f'90_runs/{run_id}'))

        deleted = self.service.delete(project_id)

        self.assertEqual(project, deleted)
        self.assertIsNone(self.repo.get(project_id))
        self.assertEqual((), runs.list_for_project(project_id))
        self.assertFalse(project_root.exists())

    def test_delete_refuses_a_project_with_an_active_run(self):
        project_id = 'c' * 32
        project = Project(project_id, 'active project', f'projects/{project_id}')
        self.repo.add(project)
        project_root = self.paths.root / project.relative_path
        project_root.mkdir(parents=True)
        run_id = 'd' * 32
        runs = RunRepository(Database(self.paths))
        runs.add(Run(run_id, project_id, Provider.CODEX, RunStatus.RUNNING,
                     f'90_runs/{run_id}'))

        with patch.object(Path, 'rename', side_effect=AssertionError(
                'active projects must be rejected before their folder is staged')) as rename:
            with self.assertRaises(ValueError):
                self.service.delete(project_id)
            rename.assert_not_called()

        self.assertEqual(project, self.repo.get(project_id))
        self.assertEqual((run_id,), tuple(run.run_id for run in runs.list_for_project(project_id)))
        self.assertTrue(project_root.is_dir())

    def test_delete_refuses_a_project_with_a_created_run(self):
        runs = RunRepository(Database(self.paths))
        project_id = 'e' * 32
        project = Project(project_id, 'created project', f'projects/{project_id}')
        self.repo.add(project)
        project_root = self.paths.root / project.relative_path
        project_root.mkdir(parents=True)
        run_id = '1' * 32
        runs.add(Run(run_id, project_id, Provider.CODEX, RunStatus.CREATED,
                     f'90_runs/{run_id}'))

        with self.assertRaises(ValueError):
            self.service.delete(project_id)

        self.assertEqual(project, self.repo.get(project_id))
        self.assertEqual((run_id,), tuple(run.run_id for run in runs.list_for_project(project_id)))
        self.assertTrue(project_root.is_dir())

    def test_delete_removes_a_paused_run_with_the_project(self):
        project_id = 'f' * 32
        project = Project(project_id, 'paused project', f'projects/{project_id}')
        self.repo.add(project)
        project_root = self.paths.root / project.relative_path
        project_root.mkdir(parents=True)
        run_id = '2' * 32
        runs = RunRepository(Database(self.paths))
        runs.add(Run(run_id, project_id, Provider.CODEX, RunStatus.PAUSED,
                     f'90_runs/{run_id}'))

        self.service.delete(project_id)

        self.assertIsNone(self.repo.get(project_id))
        self.assertEqual((), runs.list_for_project(project_id))
        self.assertFalse(project_root.exists())
    def test_copy_hash_manifest_and_opaque_layout(self):
        src = self.source()
        template = self.source('template.docx', b'docx bytes')
        project = self.service.create('中文标题', '题目\n文本', [src], template=template)
        root = self.paths.root / project.relative_path
        self.assertRegex(project.project_id, r'^[0-9a-f]{32}$')
        for name in ('00_inputs','10_analysis','20_model','30_src','40_results','50_figures','60_paper','70_review','90_runs'):
            self.assertTrue((root / name).is_dir())
        manifest = json.loads((root / '00_inputs/manifest-v0001.json').read_text('utf-8'))
        self.assertEqual(3, len(manifest['sources']))
        for entry in manifest['sources']:
            copied = root / entry['relative_path']
            self.assertEqual(entry['sha256'], hashlib.sha256(copied.read_bytes()).hexdigest())
            self.assertEqual(entry['size_bytes'], copied.stat().st_size)
        self.assertEqual(b'docx bytes', template.read_bytes())
        self.assertEqual(b'x,y\n1,2\n', src.read_bytes())
        self.assertEqual(project, self.repo.get(project.project_id))
        self.assertTrue(list((root / '10_analysis').glob('*.json')))
    def test_unique_projects(self):
        self.assertNotEqual(self.service.create('a','p',[]).project_id, self.service.create('a','p',[]).project_id)

    def test_problem_document_snapshot_is_preserved_outside_workflow_sources(self):
        for selected_name in ('原题.PDF', '原题.DOCX'):
            with self.subTest(selected_name=selected_name):
                original = b'original question snapshot\x00\xff'
                project = self.service.create(
                    'title', 'extracted question', [],
                    problem_document=(selected_name, original),
                )
                root = self.paths.root / project.relative_path
                manifest = json.loads((root / '00_inputs/manifest-v0001.json').read_text('utf-8'))
                relative = '00_inputs/problem-original-v0001' + Path(selected_name).suffix.lower()
                self.assertEqual(original, (root / relative).read_bytes())
                self.assertEqual({
                    'selected_name': selected_name, 'relative_path': relative,
                    'sha256': hashlib.sha256(original).hexdigest(), 'size_bytes': len(original),
                }, manifest['problem_document'])
                self.assertEqual(['problem'], [entry['kind'] for entry in manifest['sources']])
                self.assertEqual(project, self.repo.get(project.project_id))

    def test_invalid_problem_document_snapshot_leaves_no_committed_project(self):
        from mathmodel_paper_studio.inputs.profiling import ImportLimits
        invalid = (
            ['question.pdf', b'pdf'], ('question.pdf',), ('question.pdf', b'pdf', 'extra'),
            (Path('question.pdf'), b'pdf'), ('question.pdf', bytearray(b'pdf')),
            ('question.pdf', b''), ('question.pdf', b'12345'),
            ('folder/question.pdf', b'pdf'), (r'folder\question.pdf', b'pdf'),
            ('F:question.pdf', b'pdf'), ('question.txt', b'pdf'), ('question?.pdf', b'pdf'),
        )
        service = ProjectService(self.paths, self.repo, limits=ImportLimits(max_file_bytes=4))
        for snapshot in invalid:
            with self.subTest(snapshot=snapshot), self.assertRaises(ValueError):
                service.create('title', 'p', [], problem_document=snapshot)
            self.assertEqual((), self.repo.list_all())
            self.assertEqual([], list(self.paths.projects.iterdir()))
            self.assertEqual([], list(self.paths.temp.iterdir()))

    def test_plain_problem_import_has_no_optional_document_metadata(self):
        project = self.service.create('title', 'question', [])
        root = self.paths.root / project.relative_path
        manifest = json.loads((root / '00_inputs/manifest-v0001.json').read_text('utf-8'))
        self.assertNotIn('problem_document', manifest)
        self.assertFalse(list((root / '00_inputs').glob('problem-original-*')))

    def test_docx_data_import_preserves_optional_problem_document_metadata(self):
        from io import BytesIO
        from zipfile import ZipFile
        stream = BytesIO()
        document = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:body><w:p><w:r><w:t>observations</w:t></w:r></w:p></w:body></w:document>')
        with ZipFile(stream, 'w') as archive:
            archive.writestr('word/document.xml', document)
        source = self.source('observations.docx', stream.getvalue())
        for snapshot in (None, ('question.pdf', b'question original')):
            with self.subTest(snapshot=snapshot):
                project = self.service.create('title', 'question', [source], problem_document=snapshot)
                root = self.paths.root / project.relative_path
                manifest = json.loads((root / '00_inputs/manifest-v0001.json').read_text('utf-8'))
                self.assertEqual(['problem', 'docx_data'], [entry['kind'] for entry in manifest['sources']])
                if snapshot is None:
                    self.assertNotIn('problem_document', manifest)
                else:
                    original = manifest['problem_document']
                    self.assertEqual(snapshot[0], original['selected_name'])
                    self.assertEqual(snapshot[1], (root / original['relative_path']).read_bytes())
                profile = json.loads((root / '10_analysis/source-0001-profile.json').read_text('utf-8'))
                self.assertEqual('docx_data', profile['kind'])
                self.assertEqual(1, profile['paragraph_count'])

    def test_duplicate_and_case_variant_refused(self):
        source = self.source()
        for files in ([source,source], [source,Path(str(source).upper())]):
            with self.assertRaises(ValueError):
                self.service.create('a','p',files)
        self.assertEqual((), self.repo.list_all())
        self.assertEqual([], list(self.paths.projects.iterdir()))
    def test_partial_failure_rolls_back(self):
        good = self.source()
        bad = self.source('bad.csv', b'a,b\n1\n')
        with self.assertRaises(ValueError):
            self.service.create('a','p',[good,bad])
        self.assertEqual((), self.repo.list_all())
        self.assertEqual([], list(self.paths.projects.iterdir()))
        self.assertEqual([], list(self.paths.temp.iterdir()))
    def test_collision_and_repository_failure_preserve_existing(self):
        collision = self.paths.projects / ('a'*32)
        collision.mkdir()
        sentinel = collision / 'keep.txt'
        sentinel.write_text('keep')
        with patch('mathmodel_paper_studio.projects.project_service.uuid4') as uid:
            uid.return_value.hex = 'a'*32
            with self.assertRaises(ValueError):
                self.service.create('a','p',[])
        self.assertEqual('keep', sentinel.read_text())
        with patch.object(self.repo,'add',side_effect=RuntimeError('db unavailable')):
            with self.assertRaises(RuntimeError):
                self.service.create('b','p',[])
        self.assertEqual([collision], list(self.paths.projects.iterdir()))
    def test_junction_source_and_destination_rejected(self):
        import subprocess
        real = self.base / 'real'
        real.mkdir()
        (real / 'a.csv').write_bytes(b'x\n1\n')
        alias = self.base / 'alias'
        result = subprocess.run(['cmd','/c','mklink','/J',str(alias),str(real)],capture_output=True)
        self.assertEqual(0,result.returncode,result.stderr)
        self.addCleanup(lambda: alias.rmdir() if alias.exists() else None)
        with self.assertRaises(ValueError):
            self.service.create('a','p',[alias / 'a.csv'])
        target = self.paths.projects / ('b'*32)
        result = subprocess.run(['cmd','/c','mklink','/J',str(target),str(real)],capture_output=True)
        self.assertEqual(0,result.returncode,result.stderr)
        self.addCleanup(lambda: target.rmdir() if target.exists() else None)
        with patch('mathmodel_paper_studio.projects.project_service.uuid4') as uid:
            uid.return_value.hex = 'b'*32
            with self.assertRaises(ValueError):
                self.service.create('a','p',[])
        self.assertEqual(b'x\n1\n',(real / 'a.csv').read_bytes())

    def test_unsupported_and_duplicate_basenames(self):
        with self.assertRaises(ValueError):
            self.service.create('a', 'p', [self.source('bad.exe')])
        other = self.base / 'other'
        other.mkdir()
        same = other / 'DATA.csv'
        same.write_bytes(b'x\n1\n')
        with self.assertRaises(ValueError):
            self.service.create('a', 'p', [self.source(), same])
        self.assertEqual((), self.repo.list_all())
        self.assertEqual([], list(self.paths.temp.iterdir()))

    def test_manifest_records_and_problem_limits(self):
        from mathmodel_paper_studio.domain.models import SourceFile
        from mathmodel_paper_studio.inputs.profiling import ImportLimits
        project = self.service.create('title', '题目', [self.source()])
        root = self.paths.root / project.relative_path
        manifest = json.loads((root / '00_inputs/manifest-v0001.json').read_text('utf-8'))
        records = [SourceFile(**entry['record']) for entry in manifest['sources']]
        self.assertEqual({'problem', 'csv'}, {record.kind for record in records})
        problem = next(record for record in records if record.kind == 'problem')
        self.assertEqual('00_inputs/problem-v0001.txt', problem.relative_path)
        self.assertEqual('题目'.encode('utf-8'), (root / problem.relative_path).read_bytes())
        self.assertTrue(all('parser_notes' in entry for entry in manifest['sources']))
        self.assertNotIn(str(self.base), json.dumps(manifest))
        for value in ('', None, 'overlong'):
            with self.subTest(value=value), self.assertRaises(ValueError):
                ProjectService(self.paths, self.repo, limits=ImportLimits(max_file_bytes=3)).create('a', value, [])

    def test_project_root_junction_refused(self):
        import subprocess
        self.paths.projects.rmdir()
        outside = self.base / 'outside'
        outside.mkdir()
        (outside / 'keep').write_bytes(b'keep')
        result = subprocess.run(['cmd', '/c', 'mklink', '/J', str(self.paths.projects), str(outside)], capture_output=True)
        self.assertEqual(0, result.returncode, result.stderr)
        self.addCleanup(lambda: self.paths.projects.rmdir() if self.paths.projects.exists() else None)
        with self.assertRaises(ValueError):
            self.service.create('a', 'p', [])
        self.assertEqual([outside / 'keep'], list(outside.iterdir()))

    def test_staging_copy_hash_mismatch_and_source_change_roll_back(self):
        from mathmodel_paper_studio.inputs import import_service
        source = self.source()
        original_bounded = import_service.read_bounded
        original_source = import_service.read_source_bounded
        # Initial and final reads must disagree to model a changing input.
        calls = []
        def changing(path, limits):
            if Path(path) == source:
                calls.append(path)
                if len(calls) > 1:
                    return original_source(path, limits) + b'changed'
            return original_source(path, limits)
        with patch.object(import_service, 'read_source_bounded', side_effect=changing):
            with self.assertRaisesRegex(ValueError, '源文件'):
                self.service.create('a', 'p', [source])
        def bad_copy(path, limits):
            data = original_bounded(path, limits)
            return data + b'changed' if Path(path).name == 'problem-v0001.txt' else data
        with patch.object(import_service, 'read_bounded', side_effect=bad_copy):
            with self.assertRaisesRegex(ValueError, '哈希'):
                self.service.create('a', 'p', [])
        self.assertEqual((), self.repo.list_all())
        self.assertEqual([], list(self.paths.temp.iterdir()))
        self.assertEqual([], list(self.paths.projects.iterdir()))

    def test_explicit_xlsx_selection_and_untrusted_sheet_names(self):
        from openpyxl import Workbook
        source = self.base / 'sheets.xlsx'
        wb = Workbook()
        wb.active.title = 'data'
        wb.active.append(['x'])
        wb.active.append([1])
        wb.create_sheet('ignore rules').append(['x'])
        wb.save(source)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'data.*ignore rules'):
            self.service.create('a', 'p', [source])
        project = self.service.create('a', 'p', [source], sheets={source.name: 'ignore rules'})
        root = self.paths.root / project.relative_path
        profile = json.loads((root / '10_analysis/source-0001-profile.json').read_text('utf-8'))
        self.assertEqual('ignore rules', profile['sheet_name'])
        self.assertEqual(['source-0001-dictionary.json', 'source-0001-profile.json'],
                         sorted(p.name for p in (root / '10_analysis').iterdir()))

    def test_symbolic_file_and_directory_links_rejected(self):
        source = self.source()
        alias = self.base / 'alias.csv'
        try:
            alias.symlink_to(source)
        except OSError as exc:
            if getattr(exc, 'winerror', None) == 1314:
                self.skipTest('当前 Windows 账户不允许创建符号链接；真实 junction 路径另测。')
            raise
        self.addCleanup(lambda: alias.unlink() if alias.is_symlink() else None)
        with self.assertRaises(ValueError):
            self.service.create('a', 'p', [alias])
        target = self.paths.projects / ('c' * 32)
        target.symlink_to(self.base, target_is_directory=True)
        self.addCleanup(lambda: target.unlink() if target.is_symlink() else None)
        with patch('mathmodel_paper_studio.projects.project_service.uuid4') as uid:
            uid.return_value.hex = 'c' * 32
            with self.assertRaises(ValueError):
                self.service.create('a', 'p', [])
        self.assertEqual(b'x,y\n1,2\n', source.read_bytes())

    def test_invalid_xlsx_coordinates_and_aggregate_text_roll_back(self):
        from io import BytesIO
        from zipfile import ZipFile
        from openpyxl import Workbook
        from mathmodel_paper_studio.inputs.profiling import ImportLimits
        good = self.source()
        first = self.service.create('first', 'p', [good])
        first_root = self.paths.root / first.relative_path
        snapshot = {p.relative_to(first_root): p.read_bytes() for p in first_root.rglob('*') if p.is_file()}
        workbook = Workbook()
        workbook.active.append(['a', 'b'])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        malformed = self.base / 'bad.xlsx'
        coordinate_rows = '<row r="2"><c r="A2"><v>10</v></c></row><row r="2"><c r="A2"><v>20</v></c></row>'
        fragments = '<r><t>' + 'x' * 40 + '</t></r><r><t>' + 'y' * 40 + '</t></r>'
        rich_rows = f'<row r="2"><c r="A2" t="inlineStr"><is>{fragments}</is></c></row>'
        namespace_rows = rich_rows.replace('<r>', '<r xmlns="">')
        for rows in (coordinate_rows, rich_rows, namespace_rows):
            with self.subTest(rows=rows[:40]):
                with ZipFile(BytesIO(stream.getvalue())) as archive:
                    items = [(item, archive.read(item.filename)) for item in archive.infolist()]
                with ZipFile(malformed, 'w') as archive:
                    for item, content in items:
                        if item.filename == 'xl/worksheets/sheet1.xml':
                            content = content.replace(b'</sheetData>', rows.encode() + b'</sheetData>')
                        archive.writestr(item, content)
                before = malformed.read_bytes()
                service = ProjectService(self.paths, self.repo, limits=ImportLimits(max_cell_chars=64))
                with self.assertRaises(ValueError):
                    service.create('invalid', 'p', [good, malformed])
                self.assertEqual((first,), self.repo.list_all())
                self.assertEqual([first_root], list(self.paths.projects.iterdir()))
                self.assertEqual([], list(self.paths.temp.iterdir()))
                self.assertEqual(before, malformed.read_bytes())
                self.assertEqual(snapshot, {p.relative_to(first_root): p.read_bytes() for p in first_root.rglob('*') if p.is_file()})

    def test_update_inputs_writes_new_revision_and_preserves_old_files_and_run_outputs(self):
        original_source = self.source('original.csv', b'x,y\n1,2\n')
        project = self.service.create('title', 'old question', [original_source])
        root = self.paths.root / project.relative_path
        old_input = root / '00_inputs/source-0001-v0001.csv'
        old_input_bytes = old_input.read_bytes()
        old_manifest = root / '00_inputs/manifest-v0001.json'
        old_manifest_bytes = old_manifest.read_bytes()
        retained_run = root / '90_runs' / ('a' * 32)
        retained_run.mkdir()
        retained_artifact = retained_run / 'paper.docx'
        retained_artifact.write_bytes(b'previous run output')
        replacement = self.source('replacement.csv', b'x,y\n3,4\n')

        revision = self.service.update_inputs(
            project, 'edited question', [replacement], expected_revision=1)

        self.assertEqual(2, revision)
        manifest = json.loads((root / '00_inputs/manifest-v0002.json').read_text('utf-8'))
        self.assertEqual(2, manifest['version'])
        problem = root / '00_inputs/problem-v0002.txt'
        new_source = root / '00_inputs/source-0001-v0002.csv'
        self.assertEqual('edited question', problem.read_text('utf-8'))
        self.assertEqual(b'x,y\n3,4\n', new_source.read_bytes())
        self.assertEqual(old_input_bytes, old_input.read_bytes())
        self.assertEqual(old_manifest_bytes, old_manifest.read_bytes())
        self.assertEqual(b'previous run output', retained_artifact.read_bytes())
        self.assertEqual(hashlib.sha256(new_source.read_bytes()).hexdigest(),
                         manifest['sources'][1]['sha256'])
        self.assertEqual(project, self.repo.get(project.project_id))

    def test_update_inputs_retains_and_reverifies_original_question_document(self):
        question_document = b'original question document'
        source = self.source('original.csv', b'x,y\n1,2\n')
        project = self.service.create(
            'title', 'old question', [source],
            problem_document=('question.pdf', question_document),
        )
        root = self.paths.root / project.relative_path
        manifest_v1 = json.loads((root / '00_inputs/manifest-v0001.json').read_text('utf-8'))
        original_path = root / manifest_v1['problem_document']['relative_path']
        replacement = self.source('replacement.csv', b'x,y\n3,4\n')

        self.assertEqual(2, self.service.update_inputs(
            project, 'new question', [replacement], expected_revision=1))

        manifest_v2 = json.loads((root / '00_inputs/manifest-v0002.json').read_text('utf-8'))
        self.assertEqual(manifest_v1['problem_document'], manifest_v2['problem_document'])
        self.assertEqual(question_document, original_path.read_bytes())

    def test_update_inputs_rejects_stale_manifest_revision(self):
        project = self.service.create('title', 'question', [self.source()])
        self.service.update_inputs(project, 'first edit', [self.source('first.csv')],
                                   expected_revision=1)

        with self.assertRaisesRegex(ValueError, '版本|变化|重新读取'):
            self.service.update_inputs(project, 'stale edit', [self.source('stale.csv')],
                                       expected_revision=1)

        root = self.paths.root / project.relative_path
        self.assertTrue((root / '00_inputs/manifest-v0002.json').is_file())
        self.assertFalse((root / '00_inputs/manifest-v0003.json').exists())

    def test_concurrent_edits_from_same_revision_are_serialized(self):
        project = self.service.create('title', 'question', [self.source()])
        first_source = self.source('first.csv', b'x,y\n2,3\n')
        second_source = self.source('second.csv', b'x,y\n4,5\n')
        root = self.paths.root / project.relative_path
        manifest_entered = threading.Event()
        release_manifest = threading.Event()
        second_lock_waiting = threading.Event()
        lock_factory = project_service_module._project_edit_lock
        request_guard = threading.Lock()
        lock_requests = 0
        original_rename = Path.rename

        class TrackedLock:
            def __init__(self, lock):
                self.lock = lock

            def __enter__(self):
                nonlocal lock_requests
                with request_guard:
                    lock_requests += 1
                    request_number = lock_requests
                if request_number == 2:
                    if not self.lock.acquire(blocking=False):
                        second_lock_waiting.set()
                        self.lock.acquire()
                else:
                    self.lock.acquire()
                return self

            def __exit__(self, *_):
                self.lock.release()

        def tracked_lock(root_path):
            return TrackedLock(lock_factory(root_path))

        def pause_v2_manifest(path, target):
            if path.name == 'manifest-v0002.json' and not manifest_entered.is_set():
                manifest_entered.set()
                release_manifest.wait(10)
            return original_rename(path, target)

        with patch.object(Path, 'rename', pause_v2_manifest), \
                patch.object(project_service_module, '_project_edit_lock', tracked_lock):
            with ThreadPoolExecutor(max_workers=2) as executor:
                first = executor.submit(
                    self.service.update_inputs, project, 'first edit', [first_source],
                    expected_revision=1,
                )
                self.assertTrue(manifest_entered.wait(20), 'first save did not reach manifest commit')
                second = executor.submit(
                    self.service.update_inputs, project, 'second edit', [second_source],
                    expected_revision=1,
                )
                self.assertTrue(second_lock_waiting.wait(10), 'second save did not wait for project lock')
                release_manifest.set()
                self.assertEqual(2, first.result(timeout=20))
                with self.assertRaisesRegex(ValueError, '版本|变化|重新读取'):
                    second.result(timeout=20)

        self.assertTrue((root / '00_inputs/manifest-v0002.json').is_file())
        self.assertFalse((root / '00_inputs/manifest-v0003.json').exists())

    def test_project_edit_mutex_serializes_separate_processes(self):
        root = self.paths.projects / ('d' * 32)
        helper = self.base / 'hold_project_mutex.py'
        acquired_marker = self.base / 'mutex-acquired.txt'
        release_marker = self.base / 'mutex-release.txt'
        helper.write_text(
            'from pathlib import Path\n'
            'import sys, time\n'
            'sys.path.insert(0, sys.argv[4] + "/src")\n'
            'from mathmodel_paper_studio.projects.project_service import _project_named_mutex\n'
            'root, acquired, release = map(Path, sys.argv[1:4])\n'
            'with _project_named_mutex(root):\n'
            '    acquired.write_text("locked", encoding="utf-8")\n'
            '    deadline = time.monotonic() + 15\n'
            '    while not release.exists() and time.monotonic() < deadline:\n'
            '        time.sleep(0.01)\n',
            encoding='utf-8',
        )
        process = subprocess.Popen(
            [sys.executable, str(helper), str(root), str(acquired_marker),
             str(release_marker), str(Path(__file__).parents[3])],
            cwd=self.base, env=os.environ.copy(), stdout=subprocess.PIPE,
            stderr=subprocess.PIPE, text=True,
        )
        try:
            deadline = time.monotonic() + 10
            while not acquired_marker.exists() and process.poll() is None and time.monotonic() < deadline:
                time.sleep(0.01)
            if not acquired_marker.exists():
                self.fail(f'child failed to acquire mutex: {process.communicate(timeout=2)}')

            attempted = threading.Event()
            acquired = threading.Event()

            def acquire_from_parent():
                attempted.set()
                with project_service_module._project_named_mutex(root):
                    acquired.set()

            contender = threading.Thread(target=acquire_from_parent)
            contender.start()
            self.assertTrue(attempted.wait(5))
            self.assertFalse(acquired.wait(0.5), 'parent entered while child owns the project mutex')
            release_marker.write_text('release', encoding='utf-8')
            self.assertTrue(acquired.wait(10))
            contender.join(timeout=5)
            self.assertFalse(contender.is_alive())
        finally:
            release_marker.write_text('release', encoding='utf-8')
            if process.poll() is None:
                process.wait(timeout=10)

    def test_update_inputs_rejects_problem_text_that_main_window_cannot_reload(self):
        project = self.service.create('title', 'question', [self.source()])
        root = self.paths.root / project.relative_path
        before = (root / '00_inputs/manifest-v0001.json').read_bytes()

        with self.assertRaisesRegex(ValueError, '题目文本.*1 MiB'):
            self.service.update_inputs(
                project, 'x' * (1024 * 1024 + 1), [self.source()], expected_revision=1)

        self.assertEqual(before, (root / '00_inputs/manifest-v0001.json').read_bytes())
        self.assertFalse((root / '00_inputs/manifest-v0002.json').exists())

    def test_update_inputs_preserves_selected_xlsx_sheet(self):
        from openpyxl import Workbook
        workbook_path = self.base / 'multi-sheet.xlsx'
        workbook = Workbook()
        workbook.active.title = 'chosen'
        workbook.active.append(['x'])
        workbook.active.append([1])
        workbook.create_sheet('other').append(['ignored'])
        workbook.save(workbook_path)
        workbook.close()
        project = self.service.create('title', 'question', [workbook_path],
                                      sheets={workbook_path.name: 'chosen'})
        root = self.paths.root / project.relative_path
        copied = root / '00_inputs/source-0001-v0001.xlsx'

        self.assertEqual(2, self.service.update_inputs(
            project, 'edited question', [copied], expected_revision=1))

        profile = json.loads((root / '10_analysis/source-0001-v0002-profile.json').read_text('utf-8'))
        self.assertEqual('chosen', profile['sheet_name'])

    def test_latest_manifest_accepts_question_data_and_template_boundary(self):
        sources = [self.source(f'data-{index:02d}.csv') for index in range(32)]
        template = self.source('template.docx', b'docx template bytes')
        project = self.service.create('title', 'question', sources, template=template)

        revision, _path, manifest = latest_manifest(self.paths.root / project.relative_path)

        self.assertEqual(1, revision)
        self.assertEqual(34, len(manifest['sources']))

    def test_invalid_update_leaves_current_revision_and_project_unchanged(self):
        project = self.service.create('title', 'question', [self.source()])
        root = self.paths.root / project.relative_path
        before = {path.relative_to(root): path.read_bytes()
                  for path in root.rglob('*') if path.is_file()}
        malformed = self.source('malformed.csv', b'x,y\n1\n')

        with self.assertRaises(ValueError):
            self.service.update_inputs(project, 'edited', [malformed], expected_revision=1)

        self.assertFalse((root / '00_inputs/manifest-v0002.json').exists())
        self.assertEqual(before, {path.relative_to(root): path.read_bytes()
                                  for path in root.rglob('*') if path.is_file()})
        self.assertEqual([], list(self.paths.temp.iterdir()))

    def test_update_inputs_can_retain_current_file_and_add_another(self):
        project = self.service.create('title', 'question', [self.source()])
        root = self.paths.root / project.relative_path
        existing = root / '00_inputs/source-0001-v0001.csv'
        added = self.source('added.csv', b'a,b\n5,6\n')

        revision = self.service.update_inputs(project, 'question v2', [existing, added],
                                              expected_revision=1)

        self.assertEqual(2, revision)
        manifest = json.loads((root / '00_inputs/manifest-v0002.json').read_text('utf-8'))
        self.assertEqual(['00_inputs/problem-v0002.txt',
                          '00_inputs/source-0001-v0002.csv',
                          '00_inputs/source-0002-v0002.csv'],
                         [entry['relative_path'] for entry in manifest['sources']])
        self.assertEqual(b'x,y\n1,2\n',
                         (root / '00_inputs/source-0001-v0002.csv').read_bytes())
        self.assertEqual(b'a,b\n5,6\n',
                         (root / '00_inputs/source-0002-v0002.csv').read_bytes())
