import copy
import json
from types import SimpleNamespace
from zipfile import ZipFile

import pytest

from mathmodel_paper_studio.paper.docx_writer import write_generic_template, write_paper
from mathmodel_paper_studio.workflow.live_engine import (
    LiveWorkflowEngine, _response_format_reason,
)
from tests.fixtures.docx.fixture_factory import PNG_DATA
from tests.unit.workflow.test_paper_optional_evidence import draft_with_evidence_field


@pytest.fixture(autouse=True)
def _fake_flowchart_render(monkeypatch):
    monkeypatch.setattr(
        'mathmodel_paper_studio.workflow.live_engine.render_flowchart_png',
        lambda spec: PNG_DATA)


def result_workflow(tmp_path, *, tables=None, figures=None, mapped=False,
                    figure_evidence=None, required_figure_names=()):
    engine, context, _, _ = draft_with_evidence_field(tmp_path, 'subsection')
    evidence = (
        SimpleNamespace(evidence_id='ev_table_a', kind='table', relative_path='a.csv'),
        *(figure_evidence or (
            SimpleNamespace(evidence_id='ev_figure_a', kind='figure', relative_path='a.png'),
        )),
        SimpleNamespace(evidence_id='ev_table_b', kind='table', relative_path='b.csv'),
    )
    engine._result_bytes = lambda context, item: (
        PNG_DATA if item.kind == 'figure' else b'item,value\na,1\n')
    build_response = engine._request
    calls, originals, extras, schemas = [], [], [], []

    def generate(context, prompt, cancel_event=None):
        schema = json.loads(prompt.split('JSON schema example: ', 1)[1].splitlines()[0])
        extra = json.loads(prompt.split('<UNTRUSTED_STAGE_DATA>\n', 1)[1]
                           .split('\n</UNTRUSTED_STAGE_DATA>', 1)[0])
        batch = extra['paper_batch']['batch_id']
        calls.append(batch)
        extras.append(copy.deepcopy(extra))
        schemas.append(schema)
        value = build_response(context, set(schema), extra, cancel_event, None,
                               schema_override=schema, response_validator=lambda value: None)
        if batch == 'results':
            value['tables'] = copy.deepcopy(tables if tables is not None else [
                {'section_key': 'results', 'evidence_id': 'ev_table_b', 'caption': '边界条件审计'},
                {'section_key': 'results', 'evidence_id': 'ev_table_a', 'caption': '输入数据审计'},
            ])
            value['figures'] = copy.deepcopy(figures if figures is not None else [
                {'section_key': 'results', 'evidence_id': 'ev_figure_a', 'caption': '边界输入分布'},
            ])
        originals.append(copy.deepcopy(value))
        return json.dumps(value, ensure_ascii=False)

    engine.provider_client = SimpleNamespace(generate=generate)
    engine._request = LiveWorkflowEngine._request.__get__(engine)
    from mathmodel_paper_studio.paper.standards import get_default_paper_standard
    standard = get_default_paper_standard()

    def draft_batches():
        return engine._draft_paper_in_batches(
            context, standard, standard.required_sections + standard.optional_sections,
            set(standard.required_sections), evidence, {
                'paper_contract': {
                    'required_figure_output_names': required_figure_names,
                },
            },
            {'custom_table'} if mapped else None,
            {'custom_figure'} if mapped else None,
            None, None, tmp_path / 'result-progress.json', None,
        )

    return engine, context, evidence, draft_batches, calls, originals, extras, schemas


def test_result_parent_headings_get_distinct_word_anchors_and_resume_without_requests(tmp_path):
    engine, context, evidence, run, calls, originals, _, _ = result_workflow(tmp_path)
    combined = run()
    assert calls == ['opening', 'analysis', 'definitions', 'model_core',
                     'model_derivation', 'solution', 'results', 'closing']
    assert [item['section_key'] for item in combined['tables']] == [
        'results_table_2', 'results_table_1']
    assert combined['figures'][0]['section_key'] == 'results_figure_1'
    draft = engine._paper_draft(context, combined, evidence)
    assert [item.evidence_ids for item in draft.tables] == [('ev_table_b',), ('ev_table_a',)]
    template = write_generic_template(tmp_path, draft)
    result = write_paper(template[0], tmp_path, template[1], template[2], draft)
    with ZipFile(result.path) as archive:
        document = archive.read('word/document.xml').decode('utf-8')
        assert all(caption in document for caption in ('边界条件审计', '输入数据审计', '边界输入分布'))
        media = [name for name in archive.namelist() if name.startswith('word/media/')]
        flowcharts = sum(1 for item in draft.figures if item.visual_type == 'flowchart')
        assert len(media) == 1 + flowcharts
    cached = json.loads((tmp_path / 'result-progress.json').read_text('utf-8'))
    assert cached['batches']['results']['response'] == originals[6]
    assert cached['batches']['results']['response']['tables'][0]['section_key'] == 'results'
    assert run() == combined
    assert len(calls) == 8


def test_results_request_distinguishes_prose_keys_from_artifact_anchor_keys(tmp_path):
    explicit_tables = [{'section_key': 'results_table_7', 'evidence_id': 'ev_table_a', 'caption': '输入审计'}]
    explicit_figures = [{'section_key': 'results_figure_8', 'evidence_id': 'ev_figure_a', 'caption': '输入图'}]
    _, _, _, run, _, _, extras, schemas = result_workflow(
        tmp_path, tables=explicit_tables, figures=explicit_figures)
    combined = run()
    result_contract = extras[6]['paper_batch']
    assert result_contract['allowed_table_keys'] == ['results_table_1', 'results_table_2']
    assert result_contract['allowed_figure_keys'] == ['results_figure_1']
    assert result_contract['result_artifact_anchors'] == [
        {'kind': 'table', 'evidence_id': 'ev_table_a', 'section_key': 'results_table_1'},
        {'kind': 'table', 'evidence_id': 'ev_table_b', 'section_key': 'results_table_2'},
        {'kind': 'figure', 'evidence_id': 'ev_figure_a', 'section_key': 'results_figure_1'},
    ]
    assert 'results_table_1' in schemas[6]['tables'][0]['section_key']
    assert 'results_figure_1' in schemas[6]['figures'][0]['section_key']
    assert all('result_artifact_anchors' not in extra['paper_batch'] for extra in extras[:6])
    assert combined['tables'] == explicit_tables
    assert combined['figures'] == explicit_figures


def test_every_planned_input_figure_is_required_in_paper_results(tmp_path):
    figure_evidence = tuple(
        SimpleNamespace(evidence_id=f'ev_figure_{index}', kind='figure',
                        relative_path=f'input-series-{index:03d}.png')
        for index in range(1, 4)
    )
    figures = [
        {'section_key': 'results', 'evidence_id': item.evidence_id,
         'caption': f'观测变量 {index}'}
        for index, item in enumerate(figure_evidence, start=1)
    ]
    _, _, _, run, _, _, extras, _ = result_workflow(
        tmp_path, figures=figures, figure_evidence=figure_evidence,
        required_figure_names=tuple(item.relative_path for item in figure_evidence))

    combined = run()

    assert len(combined['figures']) == 3
    assert extras[6]['paper_batch']['required_input_figure_evidence_ids'] == [
        item.evidence_id for item in figure_evidence
    ]


def test_paper_results_reject_omitted_planned_input_figure(tmp_path):
    figure_evidence = tuple(
        SimpleNamespace(evidence_id=f'ev_figure_{index}', kind='figure',
                        relative_path=f'input-series-{index:03d}.png')
        for index in range(1, 4)
    )
    figures = [
        {'section_key': f'results_figure_{index}', 'evidence_id': item.evidence_id,
         'caption': f'观测变量 {index}'}
        for index, item in enumerate(figure_evidence[:2], start=1)
    ]
    _, _, _, run, calls, _, _, _ = result_workflow(
        tmp_path, figures=figures, figure_evidence=figure_evidence,
        required_figure_names=tuple(item.relative_path for item in figure_evidence))

    with pytest.raises(ValueError):
        run()

    assert calls[-2:] == ['results', 'results']


def test_multiple_planned_figures_can_share_one_word_heading_anchor(tmp_path):
    figure_evidence = tuple(
        SimpleNamespace(evidence_id=f'ev_figure_{index}', kind='figure',
                        relative_path=f'input-series-{index:03d}.png')
        for index in range(1, 4)
    )
    figures = [
        {'section_key': 'results_figure_1', 'evidence_id': item.evidence_id,
         'caption': f'观测变量 {index}'}
        for index, item in enumerate(figure_evidence, start=1)
    ]
    engine, context, evidence, run, _, _, _, _ = result_workflow(
        tmp_path, tables=[], figures=figures, figure_evidence=figure_evidence,
        required_figure_names=tuple(item.relative_path for item in figure_evidence))

    combined = run()
    draft = engine._paper_draft(context, combined, evidence)
    template = write_generic_template(tmp_path, draft)
    result = write_paper(template[0], tmp_path, template[1], template[2], draft)

    assert [item.section_key for item in draft.figures if item.visual_type == 'result'] == ['results_figure_1'] * 3
    with ZipFile(result.path) as archive:
        media = [name for name in archive.namelist()
                 if name.startswith('word/media/mmps-figure-')]
        flowcharts = sum(1 for item in draft.figures if item.visual_type == 'flowchart')
        assert len(media) == 3 + flowcharts
    from docx import Document
    headings = [paragraph.text for paragraph in Document(result.path).paragraphs
                if paragraph.style.name == 'Heading 3']
    assert headings == ['结果图 1']


@pytest.mark.parametrize('changes', [
    {'section_key': 'unknown_anchor'}, {'section_key': None}, {'section_key': ['results']},
    {'evidence_id': 'ev_unknown'}, {'evidence_id': 'ev_figure_a'},
    {'evidence_id': None}, {'evidence_id': ['ev_table_a']}, {'extra': 'unexpected'},
])
def test_invalid_result_keys_or_evidence_still_fail(tmp_path, changes):
    item = {'section_key': 'results', 'evidence_id': 'ev_table_a', 'caption': '输入审计', **changes}
    _, _, _, run, _, _, _, _ = result_workflow(tmp_path, tables=[item], figures=[])
    with pytest.raises(ValueError):
        run()


@pytest.mark.parametrize('second', [
    {'section_key': 'results', 'evidence_id': 'ev_table_a', 'caption': '重复结果'},
    {'section_key': 'results_table_1', 'evidence_id': 'ev_table_b', 'caption': '锚点碰撞'},
])
def test_duplicate_alias_or_explicit_anchor_collision_is_not_hidden(tmp_path, second):
    _, _, _, run, _, _, _, _ = result_workflow(tmp_path, tables=[
        {'section_key': 'results', 'evidence_id': 'ev_table_a', 'caption': '输入审计'}, second,
    ], figures=[])
    with pytest.raises(ValueError):
        run()


def test_confirmed_template_mapping_does_not_receive_builtin_alias_repair(tmp_path):
    _, _, _, run, _, _, _, _ = result_workflow(tmp_path, mapped=True)
    with pytest.raises(ValueError):
        run()


@pytest.mark.parametrize('message', [
    '图表引用了未定义或未映射的锚点键。', '表格引用了未定义或未映射的锚点键。',
    '图片引用了未定义或未映射的标题键。',
])
def test_artifact_anchor_error_is_not_repaired_as_a_prose_section_key(message):
    assert _response_format_reason(ValueError(message)) == 'result_artifacts'
